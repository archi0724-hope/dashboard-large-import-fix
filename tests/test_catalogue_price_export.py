from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from exports import catalogue_price_workbook_bytes


def test_catalogue_price_export_creates_company_tabs_without_mutating_records():
    records = pd.DataFrame([
        {"company_name": "Alpha Medical", "record_type": "price", "product_name": "Gloves", "price": 125, "source_file": "alpha-price.xlsx"},
        {"company_name": "Beta Care", "record_type": "catalogue", "product_name": "Mask", "source_file": "beta-catalogue.pdf"},
    ])

    payload = catalogue_price_workbook_bytes(records)
    book = load_workbook(BytesIO(payload), data_only=True)

    assert book.sheetnames == ["All catalogue & prices", "Alpha Medical", "Beta Care"]
    all_rows = book["All catalogue & prices"]
    assert all_rows["A5"].value == "Alpha Medical"
    assert all_rows["C6"].value == "Mask"
    assert list(records.columns) == ["company_name", "record_type", "product_name", "price", "source_file"]


def test_catalogue_price_export_uses_the_selected_columns():
    records = pd.DataFrame([{"company_name": "Alpha Medical", "product_name": "Gloves", "price": 125}])

    payload = catalogue_price_workbook_bytes(records, ["company_name", "product_name", "price"])
    sheet = load_workbook(BytesIO(payload), data_only=True)["All catalogue & prices"]

    assert [sheet.cell(4, column).value for column in range(1, 4)] == ["Company Name", "Product Name", "Price"]
