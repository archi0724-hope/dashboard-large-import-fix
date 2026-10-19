from pathlib import Path
from io import BytesIO

from openpyxl import load_workbook
import streamlit as st
from streamlit.testing.v1 import AppTest

from catalog_search import CATALOG_TEMPLATE
from storage import Store
from vendor_core import classify

APP = Path(__file__).resolve().parents[1] / "app.py"


def app(tmp_path, login=False):
    at = AppTest.from_file(str(APP), default_timeout=40)
    at.secrets["VENDOR_DATA_DIR"] = str(tmp_path)
    at.secrets["APP_PASSWORD"] = "test-admin-password"
    at.secrets["SALES_PASSWORD"] = "test-sales-password"
    if login:
        at.session_state["authenticated"] = True
    return at


def seed(tmp_path):
    store = Store(tmp_path)
    path = "Alpha Furniture/Price.csv"
    store.save_document(path, CATALOG_TEMPLATE.encode(), classify(path))
    path = "Secret Vendor/GST.txt"
    store.save_document(path, b"Private registration", classify(path))
    return store


def test_sales_login_search_and_attempted_admin_navigation(tmp_path):
    seed(tmp_path)
    at = app(tmp_path).run()
    at.text_input[0].set_value("test-sales-password")
    at.button[0].click().run()
    assert not at.exception
    assert at.session_state["access_role"] == "sales"
    assert not at.radio and not at.file_uploader
    at.chat_input(key="sales_prompt").set_value("20 office chairs under INR 3000 each; 500 5ml syringes").run()
    assert not at.exception and len(at.dataframe) == 2
    chair, syringe = [frame.value for frame in at.dataframe]
    assert chair.iloc[0]["Vendor"] == "Alpha Furniture"
    assert chair.iloc[0]["Estimated line total"] == 50000
    assert syringe.iloc[0]["Listed price"] == 4.5
    assert chair.iloc[0]["Availability"] == "Listed stock covers requested quantity"
    assert at.get("download_button")
    assert any(button.label == "Download Excel RFQ" for button in at.get("download_button"))
    at.button(key="sales_prepare_0").click().run()
    assert not at.exception
    assert any(button.label == "Download source catalog" for button in at.get("download_button"))
    at.session_state["page"] = "Data & backups"
    at.query_params["view"] = "admin"
    at.run()
    assert not at.exception and not at.radio
    assert all(button.key not in {"open_reset", "reset_history_button", "save_documents"} for button in at.button)
    assert len(Store(tmp_path).documents()) == 2
    at.button(key="sales_sign_out").click().run()
    assert not at.exception and not at.chat_input


def test_admin_setup_and_assistant_with_no_results(tmp_path):
    seed(tmp_path)
    at = app(tmp_path, login=True).run()
    at.radio(key="page").set_value("Catalog setup").run()
    assert not at.exception
    at.button(key="refresh_catalog").click().run()
    assert not at.exception
    assert at.metric[1].value == "2"
    at.radio(key="page").set_value("Sales assistant").run()
    at.chat_input(key="sales_prompt").set_value("10 oak wardrobes").run()
    assert not at.exception
    assert any("No catalog match" in item.value for item in at.info)
    assert any(button.label == "Download Excel RFQ" for button in at.get("download_button"))
    at.button(key="clear_sales_chat").click().run()
    assert not at.session_state["sales_messages"]


def test_public_sales_link_still_requires_password(tmp_path):
    at = app(tmp_path)
    at.query_params["view"] = "sales"
    at.run()
    assert not at.chat_input
    at.text_input[0].set_value("wrong-password")
    at.button[0].click().run()
    assert not at.chat_input and at.error


def test_sales_only_password_cannot_leave_admin_open(tmp_path):
    at = app(tmp_path)
    at.secrets["APP_PASSWORD"] = ""
    at.run()
    assert at.error and not at.radio and not at.chat_input


def test_identical_sales_and_admin_passwords_are_rejected(tmp_path):
    at = app(tmp_path)
    at.secrets["SALES_PASSWORD"] = "test-admin-password"
    at.run()
    assert at.error and not at.chat_input


def test_sales_rfq_download_contains_current_client_details_and_unmatched_items(tmp_path, monkeypatch):
    seed(tmp_path)
    captured = {}
    download_button = st.download_button

    def capture_download(label, data, *args, **kwargs):
        if label == "Download Excel RFQ":
            # Exercise deferred workbook generation as an actual click would.
            data = data()
            captured["book"] = load_workbook(BytesIO(data))
            captured["filename"], captured["mime"] = args[:2]
        return download_button(label, data, *args, **kwargs)

    monkeypatch.setattr(st, "download_button", capture_download)
    at = app(tmp_path).run()
    at.text_input[0].set_value("test-sales-password")
    at.button[0].click().run()
    at.chat_input(key="sales_prompt").set_value("20 office chairs; 3 oxygen concentrators").run()
    assert not at.exception
    next(item for item in at.text_input if item.label == "Client / hospital").set_value("Example hospital").run()
    next(item for item in at.text_input if item.label == "Sales executive").set_value("Example executive").run()
    assert not at.exception
    assert captured["filename"] == "Hospkart_Client_RFQ.xlsx"
    assert captured["mime"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    book = captured["book"]
    details = dict(book["RFQ details"].iter_rows(min_row=5, values_only=True))
    assert details["Client / hospital"] == "Example hospital"
    assert details["Sales executive"] == "Example executive"
    assert book["RFQ items"].max_row == 6
    all_values = " ".join(str(cell.value) for sheet in book for row in sheet for cell in row)
    assert "No catalog match" in all_values
    assert "Secret Vendor" not in all_values and "Private registration" not in all_values
    at.button(key="clear_sales_chat").click().run()
    at.chat_input(key="sales_prompt").set_value("10 office chairs").run()
    assert not at.exception
    details = dict(captured["book"]["RFQ details"].iter_rows(min_row=5, values_only=True))
    assert details["Client / hospital"] == "To be completed"
