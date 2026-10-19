import csv
from datetime import date, timedelta
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
import pytest

from catalog_search import (availability_status, extract_catalog, find_offers, gst_rate,
                            parse_requirements, rfq_rows)
from exports import rfq_workbook_bytes


def entry(**fields):
    record = {"Product": "Office chair", "Price": 100, "Currency": "INR", "Unit": "each",
              "Stock": 50, "HSN Code": "009401", "GST (%)": "18%", "Tax basis": "Exclusive",
              "Email": "sales@example.com", "Phone": "0012345678", "Contact": "Sales desk", **fields}
    output = StringIO()
    writer = csv.DictWriter(output, fieldnames=record)
    writer.writeheader()
    writer.writerow(record)
    entries, notes = extract_catalog("Price.csv", output.getvalue().encode())
    assert not notes
    return {**entries[0], "vendor": "Example vendor", "filename": "Price.csv", "document_id": 1,
            "uploaded_at": "2026-01-01T00:00:00+00:00"}


def rows_for(value, prompt="20 office chairs"):
    requirement = parse_requirements(prompt)[0]
    return rfq_rows(requirement, find_offers([value], requirement), 1)


def sheet_records(sheet):
    headers = [cell.value for cell in sheet[4]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=5, values_only=True)]


def test_csv_preserves_tax_codes_contacts_and_rfq_terms():
    value = entry(**{"Lead time": "5 days", "Payment terms": "Net 30", "Delivery terms": "Freight extra",
                     "Warranty": "1 year", "Vendor GSTIN": "GSTIN-FIXTURE"})
    assert value["gst"] == 18 and value["hsn"] == "009401"
    assert value["tax_basis"] == "Exclusive"
    assert (value["contact"], value["email"], value["phone"]) == ("Sales desk", "sales@example.com", "0012345678")
    row = rows_for(value)[0]
    assert row["Lead time"] == "5 days" and row["Payment terms"] == "Net 30"
    assert row["Warranty"] == "1 year" and row["Vendor GSTIN"] == "GSTIN-FIXTURE"
    assert row["Delivery / freight terms"] == "Freight extra"


def test_xlsx_formatted_percentage_and_hsn_keep_their_meaning():
    book = Workbook()
    sheet = book.active
    sheet.append(["Product", "Price", "Currency", "Unit", "GST (%)", "HSN Code", "GST included"])
    sheet.append(["Chair", 118, "INR", "each", 0.18, 9401, "Yes"])
    sheet["E2"].number_format = "0.00%"
    sheet["F2"].number_format = "000000"
    output = BytesIO()
    book.save(output)
    values, notes = extract_catalog("Price.xlsx", output.getvalue())
    assert not notes
    assert values[0]["gst"] == 18
    assert values[0]["hsn"] == "009401"
    assert values[0]["tax_basis"] == "Inclusive"


@pytest.mark.parametrize("raw,expected", [("18%", 18), ("5", 5), ("0%", 0), ("", None),
                                           ("extra", None), ("5-18%", None), ("18% included", None),
                                           ("101%", None), ("-5", None)])
def test_gst_unknown_is_never_silently_zero(raw, expected):
    assert gst_rate(raw) == expected


@pytest.mark.parametrize("fields,expected", [
    ({}, "Listed stock covers requested quantity"),
    ({"Stock": 3}, "Partial availability"),
    ({"Stock": 0}, "Out of stock"),
    ({"Stock": ""}, "Confirm with vendor"),
    ({"Availability": "Out of stock"}, "Out of stock"),
    ({"Stock": "", "Availability": "In stock"}, "Listed available — confirm quantity"),
    ({"Availability": "Backorder"}, "On order / request"),
    ({"Unit": "box", "Pack size": 100}, "Confirm stock unit"),
    ({"Unit": ""}, "Confirm stock unit"),
])
def test_availability_distinguishes_quantity_units_unknown_and_unavailable(fields, expected):
    requirement = parse_requirements("20 office chairs")[0]
    offers = find_offers([entry(**fields)], requirement)
    assert availability_status(offers[0], requirement).startswith(expected)


def test_old_and_conflicting_stock_cannot_look_confirmed():
    old = entry()
    newer = {**old, "stock": 0, "uploaded_at": "2026-02-01T00:00:00+00:00", "evidence": "Updated stock: 0"}
    requirement = parse_requirements("20 office chairs")[0]
    offers = find_offers([old, newer], requirement)
    old_offer = next(offer for offer in offers if offer["stock"] == 50)
    assert availability_status(old_offer, requirement).startswith("Confirm with vendor")
    conflicting = {**old, "availability": "Out of stock", "evidence": "Unavailable"}
    offers = find_offers([old, conflicting], requirement)
    assert all(not offer["eligible"] for offer in offers)
    assert all("Conflicting" in offer["status"] for offer in offers)


@pytest.mark.parametrize("price,basis,rate,net,tax,total", [
    (100, "Exclusive", "18%", 2000, 360, 2360),
    (118, "Inclusive", "18%", 2000, 360, 2360),
    (100, "Exclusive", "0%", 2000, 0, 2000),
    (1, "Inclusive", "18%", 16.95, 3.05, 20),
])
def test_rfq_gst_estimates(price, basis, rate, net, tax, total):
    row = rows_for(entry(Price=price, **{"Tax basis": basis, "GST (%)": rate}))[0]
    assert (row["Amount before GST"], row["GST amount"], row["Amount including GST"]) == (net, tax, total)


@pytest.mark.parametrize("fields", [{"GST (%)": ""}, {"Tax basis": ""}, {"Price": "On request"},
                                     {"Stock": 1}, {"MOQ": 100}, {"Currency": ""},
                                     {"Valid until": (date.today() - timedelta(days=1)).isoformat()}])
def test_unconfirmed_price_or_tax_never_produces_a_tax_total(fields):
    row = rows_for(entry(**fields))[0]
    assert row.get("Amount including GST") is None and row.get("GST amount") is None


def test_different_gst_bases_are_not_ranked_as_comparable_prices():
    first = entry()
    second = {**entry(Price=99, **{"Tax basis": "Inclusive"}), "vendor": "Second vendor"}
    offers = find_offers([first, second], parse_requirements("20 office chairs")[0])
    assert len(offers) == 2 and all(not offer["lowest"] for offer in offers)


def test_workbook_is_numeric_formula_safe_and_retains_unknown_and_unmatched_items():
    value = entry(Contact="=HYPERLINK(\"https://example.com\")", **{"GST (%)": ""})
    value["vendor"] = "=1+1"
    rows = rows_for(value)
    requirement = parse_requirements("5 oxygen concentrators")[0]
    rows.extend(rfq_rows(requirement, [], 2))
    payload = rfq_workbook_bytes(rows, {"Client / hospital": "=2+2", "Sales executive": "Test executive"})
    book = load_workbook(BytesIO(payload))
    assert book.sheetnames == ["RFQ items", "Vendor details", "RFQ details"]
    items = sheet_records(book["RFQ items"])
    assert items[0]["Listed unit price"] == 100
    assert items[0]["Requested quantity"] == 20
    assert items[0]["HSN code"] == "009401"
    assert items[0]["GST (%)"] is None and items[0]["Amount including GST"] is None
    assert items[1]["Client requirement"] == "5 oxygen concentrators"
    assert items[1]["Availability"].startswith("No catalog match")
    assert items[1]["Listed unit price"] is None
    vendors = sheet_records(book["Vendor details"])
    assert vendors[0]["Vendor phone"] == "0012345678" and vendors[0]["Source location"] == "CSV, row 2"
    details = {row["Field"]: row["Details"] for row in sheet_records(book["RFQ details"])}
    assert details["Client / hospital"] == "=2+2"
    assert details["Sales executive"] == "Test executive"
    assert "not live Hospkart" in details["Availability"]
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    assert book["RFQ items"].freeze_panes == "D5"
    assert book["RFQ items"].auto_filter.ref.endswith("6")


def test_text_catalog_and_multiple_vendor_options_export_without_inventing_fields():
    values, _ = extract_catalog("Catalog.txt", b"Office chair with mesh back")
    excerpt = {**values[0], "vendor": "Text vendor", "filename": "Catalog.txt", "uploaded_at": "2026-01-01"}
    requirement = parse_requirements("20 office chairs")[0]
    offers = find_offers([entry(), excerpt], requirement)
    rows = rfq_rows(requirement, offers, 1)
    assert [row["Vendor option"] for row in rows] == [1, 2]
    assert rows[1]["Availability"] == "Confirm with vendor"
    assert rows[1]["GST (%)"] is None and rows[1]["HSN code"] == ""
    book = load_workbook(BytesIO(rfq_workbook_bytes(rows)))
    assert len(sheet_records(book["RFQ items"])) == 2
