"""Local, source-backed catalog extraction and requirement matching.

No product, price, availability, or contact information is invented. The index is
derived from saved documents and can be rebuilt after restoring a backup.
"""
from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
import unicodedata
import zipfile
from xml.etree import ElementTree

INDEX_VERSION = 2
MAX_BYTES = 32 * 1024 * 1024
MAX_ROWS = 20000
MAX_PAGES = 250
MAX_TEXT = 2_000_000


def normalize(value) -> str:
    text = unicodedata.normalize("NFKC", str(value)).casefold()
    text = re.sub(r"(?<!\d)\.|\.(?!\d)", " ", text)
    text = re.sub(r"(\d)\s+(ml|mm|cm|kg|mg|inch|inches)\b", r"\1\2", text)
    return re.sub(r"[^\w.]+", " ", text).strip()


def number(value) -> float | None:
    """Accept one unambiguous nonnegative number, not ranges or 'from' prices."""
    text = str(value).strip()
    text = re.sub(r"(?i)\b(?:INR|USD|EUR|GBP|AED|RS)\.?\s*|[₹$€£]", "", text)
    text = text.removesuffix("/-").strip()
    if not re.fullmatch(r"(?:\d+|\d{1,3}(?:,\d{2,3})+)(?:\.\d+)?", text):
        return None
    try:
        parsed = Decimal(text.replace(",", ""))
        return float(parsed) if parsed.is_finite() and 0 <= parsed <= Decimal("1e12") else None
    except InvalidOperation:
        return None


def currency(value) -> str:
    text = str(value).upper()
    codes = set(re.findall(r"\b(?:INR|USD|EUR|GBP|AED)\b", text))
    for pattern, code in [(r"₹|\bRS\.?", "INR"), (r"€", "EUR"), (r"£", "GBP")]:
        if re.search(pattern, text):
            codes.add(code)
    # A bare $ does not establish whether the currency is USD, AUD, CAD, etc.
    return next(iter(codes)) if len(codes) == 1 else ""


def selling_unit(value, pack_size="") -> str:
    text = normalize(value)
    if text in {"pc", "pcs", "piece", "pieces", "each", "unit", "units", "nos", "no"}:
        return "each"
    if text in {"box", "boxes", "pack", "packs", "packet", "packets"}:
        size = number(pack_size)
        return f"pack of {size:g}" if size and size.is_integer() else ""
    match = re.fullmatch(r"(?:box|pack|packet)(?: of)? (\d+)", text)
    if match:
        return f"pack of {int(match[1])}"
    return text if text in {"pair", "set", "kg", "metre", "meter", "litre", "liter"} else ""


HEADERS = {
    "product": {"product", "product name", "item", "item name", "description", "item description", "product description"},
    "price": {"price", "unit price", "rate", "unit rate", "selling price", "price inr", "price rs", "unit price inr", "mrp"},
    "currency": {"currency", "currency code"},
    "unit": {"unit", "uom", "selling unit", "price unit", "unit of measure"},
    "pack_size": {"pack size", "units per pack", "pieces per box"},
    "specifications": {"specification", "specifications", "specs", "size"},
    "category": {"category", "product category"},
    "model": {"model", "sku", "item code", "product code"},
    "brand": {"brand", "make"},
    "moq": {"moq", "minimum order", "minimum order quantity"},
    "stock": {"stock", "stock quantity", "available quantity"},
    "availability": {"availability", "stock status"},
    "contact": {"contact", "phone", "email", "contact details", "vendor contact"},
    "valid_until": {"valid until", "valid till", "expiry date", "price valid until"},
    "terms": {"terms", "tax delivery terms", "tax and delivery", "price terms"},
}


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value).strip()[:2000]


def tabular_entries(rows, label: str, notes: list) -> list[dict]:
    columns = {}
    price_heading = ""
    entries = []
    for row_number, cells in enumerate(rows, 1):
        if row_number > MAX_ROWS:
            notes.append(f"{label}: stopped at {MAX_ROWS:,} rows; split this file to index the remainder.")
            break
        cells = [cell_text(value) for value in cells]
        if not columns:
            found = {key: index for index, heading in enumerate(cells)
                     for key, aliases in HEADERS.items() if normalize(heading) in aliases}
            if "product" in found:
                columns = found
                price_heading = cells[found["price"]] if "price" in found else ""
            elif row_number >= 20:
                break
            continue
        values = {key: cells[index] if index < len(cells) else "" for key, index in columns.items()}
        if not values.get("product"):
            continue
        raw_price = values.get("price", "")
        parsed_price = number(raw_price)
        entry = {key: values.get(key, "") for key in HEADERS if key != "price"}
        entry.update(price=parsed_price if parsed_price and parsed_price > 0 else None,
                     currency=currency(" ".join([values.get("currency", ""), raw_price, price_heading])),
                     unit=selling_unit(values.get("unit", ""), values.get("pack_size", "")),
                     unit_label=values.get("unit", ""),
                     moq=number(values.get("moq", "")), stock=number(values.get("stock", "")),
                     location=f"{label}, row {row_number}", kind="product",
                     evidence=" | ".join(f"{key}: {value}" for key, value in values.items() if value),
                     price_note="MRP" if normalize(price_heading) == "mrp" else "Listed price")
        entries.append(entry)
    if not columns:
        notes.append(f"{label}: no Product / Item / Description header in the first 20 rows. Use the price-list template.")
    return entries


def text_entries(text: str, location: str) -> list[dict]:
    """Keep text as evidence, never infer a price from an adjacent product row."""
    entries = []
    # Small paragraphs preserve specifications while preventing page-wide AND matches.
    for paragraph in re.split(r"\n\s*\n", text):
        lines = [line.strip() for line in paragraph.splitlines() if line.strip()]
        for start in range(0, len(lines), 6):
            excerpt = "\n".join(lines[start:start + 6])[:1800]
            if len(excerpt) < 4:
                continue
            entries.append(dict(product=lines[start][:180], specifications="", category="", model="", brand="",
                                price=None, currency="", unit="", unit_label="", moq=None, stock=None,
                                availability="", contact="", valid_until="", terms="", price_note="Quote required",
                                location=location, kind="excerpt", evidence=excerpt))
    # Only an explicit price AND selling unit on the same line establishes an
    # offer. Bare numbers in PDF tables may be model codes, sizes or pack counts.
    price_pattern = re.compile(
        r"^(?P<product>[^|\t]{3,180}?)\s*(?:\||\t|\s{2,}|\s[-–:]\s|\s+(?:price|rate)\s*:?\s*)"
        r"\s*(?P<price>(?:INR|USD|EUR|GBP|AED|Rs\.?|₹|€|£)\s*[\d,]+(?:\.\d+)?(?:/-)?)"
        r"\s*(?:\||/|per\s+)\s*(?P<unit>each|pcs?|pieces?|units?|nos|pair|set|(?:pack|box)\s+of\s+\d+)\s*$", re.I)
    for line_number, line in enumerate(text.splitlines(), 1):
        match = price_pattern.fullmatch(line.strip())
        if not match or not re.search(r"[a-z]", match["product"], re.I):
            continue
        price = number(match["price"])
        if price is None or price <= 0:
            continue
        entries.append(dict(product=match["product"].strip(), specifications="", category="", model="", brand="",
                            price=price, currency=currency(match["price"]), unit=selling_unit(match["unit"]),
                            unit_label=match["unit"], moq=None, stock=None, availability="", contact="",
                            valid_until="", terms="", price_note="Listed price", location=f"{location}, line {line_number}",
                            kind="product", evidence=line.strip()))
    return entries


def extract_catalog(filename: str, payload: bytes) -> tuple[list[dict], list[str]]:
    notes = []
    if len(payload) > MAX_BYTES:
        return [], ["File exceeds the 32 MB indexing limit. Split or compress it and upload again."]
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = payload.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return tabular_entries(csv.reader(StringIO(text), dialect), "CSV", notes), notes
    if suffix in {".xlsx", ".docx"}:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 128 * 1024 * 1024:
                return [], ["Expanded Office file exceeds the indexing limit. Split this file."]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        try:
            entries = []
            for sheet in workbook.worksheets[:30]:
                entries.extend(tabular_entries(sheet.iter_rows(values_only=True), sheet.title, notes))
                if len(entries) >= MAX_ROWS:
                    notes.append(f"Stopped at {MAX_ROWS:,} products; split this workbook to index all rows.")
                    break
            if len(workbook.worksheets) > 30:
                notes.append("Only the first 30 sheets were indexed.")
            return entries[:MAX_ROWS], notes
        finally:
            workbook.close()
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(payload), strict=False)
        entries, characters, empty_pages = [], 0, 0
        for index, page in enumerate(reader.pages[:MAX_PAGES], 1):
            try:
                text = page.extract_text() or ""
            except Exception:
                notes.append(f"Page {index}: text extraction failed.")
                continue
            characters += len(text)
            if characters > MAX_TEXT or len(entries) >= MAX_ROWS:
                notes.append("Text limit reached. Split this PDF to index the remaining pages.")
                break
            if not text.strip():
                empty_pages += 1
            entries.extend(text_entries(text, f"Page {index}"))
        if empty_pages:
            notes.append(f"{empty_pages} pages have no readable text; upload an OCR/text version or a structured price list.")
        if len(reader.pages) > MAX_PAGES:
            notes.append(f"Only the first {MAX_PAGES} pages were indexed.")
        return entries[:MAX_ROWS], notes
    if suffix == ".docx":
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            root = ElementTree.fromstring(archive.read("word/document.xml"))
        namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = ["".join(node.itertext()) for node in root.findall(".//w:p", namespace)]
        text = "\n\n".join(paragraphs)
    elif suffix == ".txt":
        text = payload.decode("utf-8-sig", errors="replace")
    else:
        return [], ["This format needs conversion. Upload a text PDF, DOCX, TXT, CSV or XLSX; scans need OCR first."]
    if len(text) > MAX_TEXT:
        notes.append("Text was truncated; split the document to index the rest.")
    return text_entries(text[:MAX_TEXT], "Document text")[:MAX_ROWS], notes


def catalog_documents(documents):
    """Only assigned commercial documents may be exposed to sales."""
    commercial_types = {"Catalogue", "Price"}
    flags = documents.types.map(lambda values: bool(set(values) & commercial_types) and set(values) <= commercial_types)
    return documents[flags & documents.available & documents.company_key.ne("")].copy()


def refresh_catalog(store, documents, progress=None) -> dict:
    index = store.catalog_index()
    candidates = catalog_documents(documents)
    updated = 0
    for position, doc in enumerate(candidates.itertuples(index=False), 1):
        existing = index.get(int(doc.id))
        if existing and existing["file_hash"] == doc.file_hash and existing["version"] == INDEX_VERSION:
            continue
        if progress:
            progress(position, len(candidates), doc.filename)
        try:
            if doc.size_bytes > MAX_BYTES:
                entries, notes = [], ["File exceeds the 32 MB indexing limit. Split or compress it."]
            else:
                payload = store.read_bytes(doc.id)
                if not payload:
                    continue
                entries, notes = extract_catalog(doc.filename, payload)
        except Exception:
            # Do not cache transient read/parser errors; the next refresh can retry.
            continue
        store.save_catalog_index(int(doc.id), doc.file_hash, INDEX_VERSION, entries, notes)
        updated += 1
    return {"updated": updated, "catalogs": len(candidates)}


def load_catalog(store, documents) -> tuple[list[dict], list[dict], int]:
    index, entries, reports, pending = store.catalog_index(), [], [], 0
    for doc in catalog_documents(documents).itertuples(index=False):
        indexed = index.get(int(doc.id))
        if not indexed or indexed["file_hash"] != doc.file_hash or indexed["version"] != INDEX_VERSION:
            pending += 1
            reports.append({"Vendor": doc.company_name, "File": doc.filename, "Search entries": 0,
                            "Status": "Pending — refresh catalog search"})
            continue
        extracted = json.loads(indexed["entries_json"])
        for entry in extracted:
            entries.append({**entry, "document_id": int(doc.id), "vendor": doc.company_name,
                            "filename": doc.filename, "uploaded_at": doc.uploaded_at,
                            "indexed_at": indexed["indexed_at"]})
        notes = json.loads(indexed["notes_json"])
        reports.append({"Vendor": doc.company_name, "File": doc.filename, "Search entries": len(extracted),
                        "Status": "; ".join(notes) if notes else ("Ready" if extracted else "No readable products or text")})
    return entries, reports, pending


SINGULAR = {"chairs": "chair", "tables": "table", "desks": "desk", "beds": "bed", "cots": "bed", "cot": "bed",
            "syringes": "syringe", "cabinets": "cabinet", "trolleys": "trolley", "stools": "stool", "sofas": "sofa",
            "needles": "needle", "gloves": "glove", "wheelchairs": "wheelchair", "cupboards": "cabinet", "cupboard": "cabinet"}
STOP_WORDS = set("i we our me my us team need needs want wants require required requirement requirements looking for find show please a an the of with and or to from in which who vendor vendors vender venders supplier suppliers has have sell sells selling supply supplies best cheapest lowest price prices cost budget buy buying purchase give get details data according available availability can you any type types furnitureprice each per under below than less upto up at is are be would like compare quotation quote list all only qty quantity units unit pcs pieces piece nos rs inr usd eur gbp aed include including provide hospitalgrade".split())
FURNITURE = {"furniture", "chair", "table", "desk", "bed", "cabinet", "trolley", "stool", "sofa", "locker"}


def tokens(text: str) -> set[str]:
    text = re.sub(r"\bss\b", "stainless steel", normalize(text))
    text = re.sub(r"\bnon\s+(\w+)", r"non\1", text)
    return {SINGULAR.get(word, word) for word in text.split() if word not in STOP_WORDS}


@dataclass
class Requirement:
    text: str
    terms: set[str]
    quantity: int = 1
    budget: float | None = None
    currency: str = ""
    unit: str = "each"


def parse_requirements(prompt: str) -> list[Requirement]:
    if not prompt.strip() or len(prompt) > 4000:
        raise ValueError("Describe your requirement in 1–4,000 characters.")
    # Commas inside numbers and 'and' inside a specification are kept intact.
    parts = re.split(r"\n|;|,(?!\d)|\s+and\s+(?=\d+\s+(?!ml\b|mm\b|cm\b))", prompt, flags=re.I)
    parts = [part.strip() for part in parts if part.strip()]
    if len(parts) > 12:
        raise ValueError("Search up to 12 requirement lines at a time.")
    result = []
    for part in parts:
        working, budget, budget_currency = part, None, ""
        match = re.search(r"(?i)\b(?:under|below|up to|upto|less than|max(?:imum)?(?: price)?|budget(?: of)?)\s*(?:of\s*)?((?:INR|USD|EUR|GBP|AED|Rs\.?|[₹$€£])?\s*[\d,]+(?:\.\d+)?)\s*(INR|USD|EUR|GBP|AED)?", working)
        if match:
            budget = number(match[1])
            budget_currency = currency(match[0])
            working = working[:match.start()] + working[match.end():]
        working = re.sub(r"^\s*[-•]\s*|^\s*\d+[.)]\s+", "", working)
        working = re.sub(r"(?i)^(?:\s*(?:i|we|are|need|want|require|find|show|please|looking|for|would|like|to|buy|can|you|me|us|get|which|vendor|has)\b)+\s*", "", working)
        quantity = 1
        quantity_match = re.search(r"(?i)\b(?:qty|quantity)\s*[:=]?\s*([\d,]+)\b", working)
        if not quantity_match:
            quantity_match = re.match(r"(?i)^\s*(?:(?:i|we)\s+)?(?:(?:need|want|require|find|show|please)\s+)*([\d,]+)\s+(?!(?:ml|mm|cm|kg|mg|inch)\b)", working)
        if quantity_match:
            quantity = int(quantity_match[1].replace(",", ""))
            working = working[:quantity_match.start()] + working[quantity_match.end():]
        if not 1 <= quantity <= 1_000_000_000:
            raise ValueError("Quantity must be between 1 and 1,000,000,000.")
        unit = "each"
        pack = re.search(r"(?i)\b(?:boxes?|packs?)(?: of)?\s+(\d+)\b", working)
        if pack:
            unit = f"pack of {int(pack[1])}"
            working = working[:pack.start()] + working[pack.end():]
        terms = tokens(working)
        if not terms:
            raise ValueError("Include a product name, for example '20 office chairs' or '500 5ml syringes'.")
        result.append(Requirement(part, terms, quantity, budget, budget_currency, unit))
    return result


def match_terms(required: set[str], available: set[str]) -> bool:
    return all(bool(FURNITURE & available) if term == "furniture" else term in available for term in required)


def find_offers(entries: list[dict], requirement: Requirement) -> list[dict]:
    def product_identity(entry):
        return tuple(normalize(entry.get(key, "")) for key in
                     ("vendor", "product", "specifications", "brand", "model", "currency", "unit"))

    # A historical cheap quote must not outrank the vendor's newer price list,
    # including when the newer price exceeds the customer's budget.
    versions = {}
    for entry in entries:
        if entry["kind"] != "product":
            continue
        identity, uploaded = product_identity(entry), entry["uploaded_at"]
        quote = tuple(entry.get(key) for key in ("price", "moq", "stock", "valid_until", "terms"))
        latest, quotes = versions.get(identity, ("", set()))
        if uploaded > latest:
            versions[identity] = (uploaded, {quote})
        elif uploaded == latest:
            quotes.add(quote)
    matches = []
    for entry in entries:
        searchable = " ".join(str(entry.get(key, "")) for key in
                              ("product", "specifications", "category", "brand", "model") if entry.get(key))
        if entry["kind"] == "excerpt":
            searchable = entry["evidence"]
        if not match_terms(requirement.terms, tokens(searchable)):
            continue
        issues = []
        price = entry.get("price")
        price_ready = price is not None and bool(entry["currency"]) and bool(entry["unit"])
        if entry["kind"] == "excerpt":
            issues.append("Catalog mention — confirm specifications and request a quote")
        elif price is None:
            issues.append("Quote required")
        elif not entry["currency"] or not entry["unit"]:
            issues.append("Confirm currency / selling unit")
        if entry["kind"] == "product":
            latest, quotes = versions[product_identity(entry)]
            if entry["uploaded_at"] < latest:
                issues.append("Older listing — a newer vendor listing is available")
            elif len(quotes) > 1:
                issues.append("Conflicting vendor listings — confirm the applicable quote")
        if entry["unit"] and entry["unit"] != requirement.unit:
            issues.append("Different selling unit — confirm quantity and price basis")
        if entry.get("moq") is not None and entry["moq"] > requirement.quantity:
            issues.append(f"Below minimum order of {entry['moq']:g}")
        if entry.get("stock") is not None and entry["stock"] < requirement.quantity:
            issues.append("Listed stock is below requested quantity")
        if normalize(entry.get("availability", "")) in {"out of stock", "unavailable", "discontinued", "no", "not available"}:
            issues.append("Listed as unavailable")
        if entry.get("valid_until"):
            try:
                if date.fromisoformat(entry["valid_until"]) < date.today():
                    issues.append("Price expired — request an updated quote")
            except ValueError:
                issues.append("Confirm price validity date")
        if requirement.budget is not None:
            if not requirement.currency:
                issues.append("Add budget currency, e.g. 'under INR 5000 each'")
            elif price_ready and entry["unit"] == requirement.unit:
                if entry["currency"] != requirement.currency:
                    issues.append("Different budget currency")
                elif price > requirement.budget:
                    continue
        eligible = price_ready and not issues
        total = float((Decimal(str(price)) * requirement.quantity).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)) if eligible else None
        # Compare like descriptions/specifications/brands and price terms only.
        basis = tuple(normalize(entry.get(key, "")) for key in
                      ("product", "specifications", "brand", "model", "currency", "unit", "terms", "price_note"))
        matches.append({**entry, "quantity": requirement.quantity, "total": total, "eligible": eligible,
                        "status": "; ".join(issues) or "Listed offer — confirm stock and final quote",
                        "comparison_basis": basis, "lowest": False})
    # Deduplicate repeated catalog passages/rows while retaining source provenance.
    unique = {}
    for entry in matches:
        key = (entry["vendor"], entry["evidence"], entry["price"], entry["currency"], entry["unit"])
        previous = unique.get(key)
        if not previous or entry["uploaded_at"] > previous["uploaded_at"]:
            unique[key] = entry
    matches = list(unique.values())
    groups = {}
    for offer in matches:
        if offer["eligible"]:
            groups.setdefault(offer["comparison_basis"], []).append(offer)
    for group in groups.values():
        if len({offer["vendor"] for offer in group}) > 1:
            lowest = min(offer["price"] for offer in group)
            for offer in group:
                offer["lowest"] = offer["price"] == lowest
    return sorted(matches, key=lambda offer: (not offer["eligible"], offer["currency"], offer["unit"],
                                              offer["price"] if offer["price"] is not None else float("inf"), offer["vendor"]))


def comparison_rows(requirement: Requirement, offers: list[dict]) -> list[dict]:
    return [{"Requirement": requirement.text, "Vendor": offer["vendor"], "Product / catalog text": offer["product"],
             "Specifications": offer.get("specifications", ""), "Brand / model": " / ".join(filter(None, [offer.get("brand"), offer.get("model")])),
             "Listed price": offer["price"], "Currency": offer["currency"] or "Not stated", "Unit": offer["unit"] or offer.get("unit_label") or "Not stated",
             "Requested quantity": requirement.quantity, "Estimated line total": offer["total"],
             "Minimum order": offer.get("moq"), "Listed stock": offer.get("stock"),
             "Valid until": offer.get("valid_until", ""), "Tax / delivery terms": offer.get("terms", ""),
             "Contact": offer.get("contact", ""), "Comparison": "Lowest comparable listed price" if offer["lowest"] else "",
             "Status": offer["status"], "Source": f"{offer['filename']} · {offer['location']}",
             "Catalog uploaded": offer["uploaded_at"][:10]} for offer in offers]


CATALOG_TEMPLATE = ("Product,Specifications,Category,Brand,Model,Price,Currency,Unit,Pack size,MOQ,Stock,Contact,Valid until,Terms\n"
                    "Office chair,Mesh back,Furniture,Example brand,CHAIR-01,2500,INR,each,,1,50,sales@example.com,2027-12-31,Tax and delivery extra\n"
                    "Syringe,5ml sterile,Consumables,Example brand,SYR-05,4.5,INR,each,,100,10000,sales@example.com,2027-12-31,Tax and delivery extra\n")
