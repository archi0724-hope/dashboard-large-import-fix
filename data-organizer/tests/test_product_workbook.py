import csv

from openpyxl import load_workbook
import pytest

from app.config import RuntimeConfig
from app.export.product_workbook import create_product_workbook, amount
from app.pipeline import LocalSource
from app.services import _company_folder_from_source_path
from .test_api_export import client


def test_company_product_workbook_preserves_fields_and_flags_uncertainty(settings, tmp_path):
    folder = tmp_path / "input"
    folder.mkdir()
    path = folder / "catalogue.csv"
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["Company Name", "Product Name", "SKU", "Price", "MRP", "Specification", "Pack Size", "Currency", "Email"])
        writer.writerows([
            ["Beta", "Wheelchair", "W2", "100-200", "250", "Adjustable back", "1 unit", "INR", "beta@example.org"],
            ["Alpha", "Gloves", "G2", "₹ 125", "150", "Sterile, powder-free", "Box of 100", "INR", "alpha@example.org"],
            ["Alpha", "Gloves", "G1", "100", "150", "Latex", "Box of 50", "INR", ""],
            ["Alpha", "Gloves", "G1", "110", "150", "Latex", "Box of 50", "INR", ""],
            ["Alpha", "=unsafe product", "F1", "", "999", "", "", "INR", ""],
        ])
    before = path.read_bytes()
    output = tmp_path / "output" / "products.xlsx"
    result = create_product_workbook(LocalSource(folder), settings, RuntimeConfig(), output, _company_folder_from_source_path)
    assert result["products"] == 5
    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Products & Prices", "Company Contacts", "Needs Review", "File Processing"]
    sheet = workbook["Products & Prices"]
    rows = list(sheet.values)
    records = [dict(zip(rows[0], row)) for row in rows[1:]]
    assert [(r["Company Name"], r["Product Name"], r["SKU"]) for r in records] == sorted(
        [(r["Company Name"], r["Product Name"], r["SKU"]) for r in records], key=lambda row: tuple(value.lower() for value in row))
    glove = next(r for r in records if r["SKU"] == "G2")
    assert glove["Price"] == 125 and glove["MRP"] == 150
    assert glove["Specification"] == "Sterile, powder-free" and glove["Pack Size"] == "Box of 100"
    uncertain = next(r for r in records if r["SKU"] == "W2")
    assert uncertain["Price"] == "100-200" and uncertain["MRP"] == 250
    missing = next(r for r in records if r["SKU"] == "F1")
    assert missing["Price"] is None and missing["MRP"] == 999
    assert sheet.cell(2, 3).data_type == "s"
    issues = list(workbook["Needs Review"].values)
    assert any("Conflicting" in str(row) for row in issues)
    assert any("Unclear numeric" in str(row) for row in issues)
    assert workbook["Company Contacts"].max_row == 3
    assert path.read_bytes() == before


@pytest.mark.parametrize("value,expected", [("₹1,250.50", 1250.5), ("100-200", "100-200"), ("on request", "on request"), (0, 0)])
def test_price_parser_does_not_choose_one_number_from_a_range(value, expected):
    assert amount(value)[0] == expected


def test_large_batch_workers_account_for_every_file(settings, tmp_path):
    folder = tmp_path / "source" / "Acme"
    folder.mkdir(parents=True)
    for index in range(17):
        (folder / f"price-{index}.csv").write_text(
            f"Product Name,SKU,Price,Currency\nGlove {index},G{index},{index+1},INR\n", encoding="utf-8")
    output = tmp_path / "large-batch.xlsx"
    result = create_product_workbook(LocalSource(tmp_path / "source"), settings, RuntimeConfig(), output,
                                     _company_folder_from_source_path)
    assert result["products"] == result["files"] == 17
    book = load_workbook(output)
    assert book["File Processing"].max_row == 18


def test_alternate_headers_keep_primary_price_and_extra_source_values(settings, tmp_path):
    folder = tmp_path / "source" / "Acme"
    folder.mkdir(parents=True)
    (folder / "prices.csv").write_text("ModelNo,Particulars,Rate Each,Rate Each\nM1,Bed,46000,400\n", encoding="utf-8")
    (folder / "reference.csv").write_text("Product Name,Price\nReference only,99\n", encoding="utf-8")
    output = tmp_path / "workbook.xlsx"
    result = create_product_workbook(LocalSource(tmp_path / "source"), settings, RuntimeConfig(), output,
                                     _company_folder_from_source_path, reference_files=["Acme/reference.csv"])
    assert result["products"] == result["files"] == 1
    book = load_workbook(output)
    values = list(book["Products & Prices"].values)
    record = dict(zip(values[0], values[1]))
    assert record["SKU"] == "M1" and record["Price"] == 46000
    assert '400' in record["Notes"] and record["Review Flag"] == "Needs review"


def test_manufacturer_is_separate_from_supplier_and_generic_tax_is_not_guessed(settings, tmp_path):
    folder = tmp_path / "source" / "Acme"
    folder.mkdir(parents=True)
    (folder / "prices.csv").write_text("ManufacturerName,ProductName,SKU,Price,Tax,Currency\nBrand Ltd,Glove,G1,100,5,INR\n", encoding="utf-8")
    output = tmp_path / "workbook.xlsx"
    create_product_workbook(LocalSource(tmp_path / "source"), settings, RuntimeConfig(), output, _company_folder_from_source_path)
    book = load_workbook(output)
    values = list(book["Products & Prices"].values)
    record = dict(zip(values[0], values[1]))
    assert record["Company Name"] == "Acme" and record["Brand"] == "Brand Ltd"
    assert record["GST"] is None and 'Tax' in record["Notes"]


def test_background_api_produces_downloadable_workbook(client):
    import time
    response = client.post("/api/export/products")
    assert response.status_code == 200
    for _ in range(100):
        job = client.get("/api/export/products/status").json()
        if not job["running"]:
            break
        time.sleep(.1)
    assert not job["running"] and not job["error"]
    download = client.get("/api/exports/download", params={"file": job["result"]["file"]})
    assert download.status_code == 200 and download.content[:2] == b"PK"


@pytest.mark.parametrize("extension", [".png", ".jpg", ".pdf"])
def test_scanned_product_rows_use_ocr_and_keep_original_files(settings, tmp_path, extension):
    from PIL import Image, ImageDraw, ImageFont
    from app.ingestion.ocr_reader import ocr_available
    if not ocr_available(settings.tesseract_cmd):
        pytest.skip("Tesseract is needed for actual scanned catalogue reading")
    folder = tmp_path / "input" / "Acme"
    folder.mkdir(parents=True)
    image = Image.new("RGB", (1600, 300), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("arial.ttf", 36)
    for x, header, value in ((50, "Product Name", "Gloves"), (800, "SKU", "G1"), (1100, "Price INR", "125")):
        draw.text((x, 40), header, fill="black", font=font)
        draw.text((x, 130), value, fill="black", font=font)
    path = folder / ("catalogue" + extension)
    image.save(path, resolution=150.0)
    before = path.read_bytes()
    output = tmp_path / "products.xlsx"
    result = create_product_workbook(LocalSource(tmp_path / "input"), settings, RuntimeConfig(ocr_enabled=True),
                                     output, _company_folder_from_source_path)
    assert result["products"] == 1
    book = load_workbook(output)
    values = list(book["Products & Prices"].values)
    record = dict(zip(values[0], values[1]))
    assert record["Company Name"] == "Acme" and record["Product Name"] == "Gloves"
    assert record["SKU"] == "G1" and record["Price"] == 125
    assert "review" in record["Extraction Confidence"].lower()
    assert book["Needs Review"].max_row > 1 and path.read_bytes() == before
