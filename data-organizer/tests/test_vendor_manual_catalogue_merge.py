import sys
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "vendor_dashboard"))

from exports import (
    CATALOGUE_PRICE_COLUMN_OPTIONS,
    MAX_MANUAL_CATALOGUE_BYTES,
    MAX_MANUAL_CATALOGUE_ROWS,
    catalogue_price_workbook_bytes,
    merge_catalogue_price_workbook_bytes,
)


def manual_workbook(sheets):
    output = BytesIO()
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(output)
    return output.getvalue()


def test_manual_rows_append_to_generated_workbook_without_changing_source_records():
    records = pd.DataFrame(
        [
            {
                "company_name": "Alpha Medical",
                "record_type": "price",
                "product_name": "Medical examination gloves",
                "sku": "G-1",
                "price": 100,
                "currency": "INR",
                "source_file": "catalogue.csv",
            },
        ]
    )
    manual = manual_workbook(
        [
            (
                "Manual prices",
                [
                    [
                        "Company Name",
                        "Record Type",
                        "Product Name",
                        "SKU",
                        "Price",
                        "Currency",
                    ],
                    ["Alpha Medical", "price", "Medical examination gloves", "G-1", 100, "INR"],
                    ["Beta Medical", "catalogue", "Wheelchair", "W-2", 5000, "INR"],
                ],
            ),
        ]
    )

    payload, added, manual_preview, review_preview, excluded = merge_catalogue_price_workbook_bytes(
        records, manual, "manual.xlsx"
    )

    assert added == 2
    assert len(manual_preview) == 2
    assert set(manual_preview["company_name"]) == {"Alpha Medical", "Beta Medical"}
    assert review_preview.empty
    assert excluded == 0
    assert len(records) == 1
    workbook = load_workbook(BytesIO(payload), data_only=True)
    all_rows = list(workbook["All catalogue & prices"].values)
    headers = all_rows[3]
    products = [dict(zip(headers, row)) for row in all_rows[4:]]
    assert [(row["Company Name"], row["Product Name"]) for row in products] == [
        ("Alpha Medical", "Medical examination gloves"),
        ("Alpha Medical", "Medical examination gloves"),
        ("Beta Medical", "Wheelchair"),
    ]
    manual_rows = [
        row for row in products if str(row["Source File"]).startswith("Manual upload:")
    ]
    assert len(manual_rows) == 2
    assert {row["Source Sheet"] for row in manual_rows} == {"Manual prices"}
    assert "Beta Medical" in workbook.sheetnames
    assert all(
        cell.data_type != "f" for sheet in workbook for row in sheet for cell in row
    )
    workbook.close()


def test_generated_company_tabs_are_not_reimported_when_all_rows_sheet_exists():
    headers = [label for label, _ in CATALOGUE_PRICE_COLUMN_OPTIONS.values()]
    original = manual_workbook(
        [
            (
                "All catalogue & prices",
                [
                    ["ALL COMPANIES"],
                    ["note"],
                    [],
                    headers,
                    [
                        "Alpha",
                        "price",
                        "Medical examination gloves",
                        "G-1",
                        "",
                        "",
                        "",
                        "",
                        "100",
                        "",
                        "",
                        "INR",
                    ],
                ],
            ),
            (
                "Alpha",
                [
                    ["Alpha"],
                    ["note"],
                    [],
                    headers,
                    [
                        "Alpha",
                        "price",
                        "Medical examination gloves",
                        "G-1",
                        "",
                        "",
                        "",
                        "",
                        "100",
                        "",
                        "",
                        "INR",
                    ],
                ],
            ),
        ]
    )

    payload, added, _, _, _ = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), original, "generated.xlsx"
    )

    assert added == 1
    workbook = load_workbook(BytesIO(payload), data_only=True)
    all_rows = list(workbook["All catalogue & prices"].values)
    headers = next(
        row for row in all_rows if "Company Name" in row and "Product Name" in row
    )
    header_index = all_rows.index(headers)
    assert (
        sum(
            any(value not in (None, "") for value in row)
            for row in all_rows[header_index + 1 :]
        )
        == 1
    )
    workbook.close()


def test_vendor_generated_workbook_can_be_merged_directly():
    records = pd.DataFrame([
        {
            "company_name": "Acme Medical",
            "record_type": "price",
            "product_name": "ZX Medical gloves",
            "sku": "ZX-1",
            "price": 125,
        },
    ])
    generated = catalogue_price_workbook_bytes(records)

    merged, added, _, _, _ = merge_catalogue_price_workbook_bytes(
        records, generated, "Company_Wise_Catalogue_and_Prices.xlsx"
    )

    assert added == 1
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert len(products) == 2
    assert all(row["Product Name"] == "ZX Medical gloves" for row in products)
    workbook.close()


def test_organizer_headers_and_company_named_worksheet_are_supported():
    manual = manual_workbook([
        ("Acme Medical", [
            ["Organizer source export"],
            [],
            ["Vendor Name", "Item Description", "Code", "Dealer Rate", "GST Percent"],
            ["Acme Medical", "Sterile gloves", "SG-10", 125, 5],
        ]),
        ("Beta Medical", [
            ["Product", "Price", "Currency"],
            ["Wheelchair", 5000, "INR"],
        ]),
        ("Company Contacts", [
            ["Company Name", "Contact Person", "Phone"],
            ["Acme Medical", "A Person", "000"],
        ]),
    ])

    merged, added, _, _, _ = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), manual, "Organizer_Products_and_Prices.xlsx"
    )

    assert added == 2
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert [(row["Company Name"], row["Product Name"]) for row in products] == [
        ("Acme Medical", "Sterile gloves"),
        ("Beta Medical", "Wheelchair"),
    ]
    assert products[0]["SKU"] == "SG-10"
    assert products[0]["Price"] == "125"
    workbook.close()


def test_vendor_product_prices_workbook_headers_are_supported():
    manual = manual_workbook([
        ("Vendor Product Prices", [
            ["Vendor Name", "Date of Conversation", "Product Names", "Product Prices"],
            [
                "Devansh Medical",
                pd.Timestamp("2026-08-05").to_pydatetime(),
                "Clot Activator Vials, Non-Vacuum (Red)",
                "₹1.40",
            ],
            [
                "Devansh Medical",
                pd.Timestamp("2026-09-02").to_pydatetime(),
                "Vacuum ESR 13x75 (Vacucare), Vacuum PT 13x75 (Vacucare)",
                "₹2.95, ₹2.95",
            ],
            ["Devansh Medical", pd.Timestamp("2026-09-03").to_pydatetime(), ", ", "₹1.00, ₹2.00"],
        ]),
    ])

    merged, added, manual_preview, review_preview, excluded = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), manual, "ARCHI CLEAN AND EXTRATION OF THE DEVANSH SIR DATA .xlsx"
    )

    assert added == 2
    assert review_preview.empty
    assert excluded == 0
    assert len(manual_preview) == 2
    assert "conversation_date" in manual_preview
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert len(products) == 2
    assert all(row["Company Name"] == "Devansh Medical" for row in products)
    assert all(row["Record Type"] == "price" for row in products)
    assert all(row["Source File"].endswith("DATA .xlsx") for row in products)
    assert {row["Date of Conversation"] for row in products} == {
        "2026-08-05 00:00:00",
        "2026-09-02 00:00:00",
    }
    assert {row["Price"] for row in products} == {"1.4", "₹2.95, ₹2.95"}
    assert {row["Product Name"] for row in products} == {
        "Clot Activator Vials, Non-Vacuum (Red)",
        "Vacuum ESR 13x75 (Vacucare), Vacuum PT 13x75 (Vacucare)",
    }
    workbook.close()


@pytest.mark.parametrize(
    "content,filename,message",
    [
        (b"not an xlsx", "manual.xlsx", "readable"),
        (b"anything", "manual.xls", ".xlsx"),
        (
            manual_workbook([("Products & Prices", [["Product Name", "Price"], ["Gloves", 125]])]),
            "manual.xlsx",
            "Company Name and Product Name",
        ),
    ],
)
def test_invalid_manual_workbooks_are_rejected(content, filename, message):
    with pytest.raises(ValueError, match=message):
        merge_catalogue_price_workbook_bytes(pd.DataFrame(), content, filename)


def test_empty_generated_workbook_reports_that_it_has_no_product_rows():
    empty_export = catalogue_price_workbook_bytes(pd.DataFrame())

    with pytest.raises(ValueError, match="contains column headers but no product rows"):
        merge_catalogue_price_workbook_bytes(
            pd.DataFrame(), empty_export, "Company Products and Prices archi sorted.xlsx"
        )


def test_empty_generated_target_can_be_merged_with_a_populated_source():
    empty_target = catalogue_price_workbook_bytes(pd.DataFrame())
    manual_source = manual_workbook([
        ("Vendor Product Prices", [
            ["Vendor Name", "Product Names", "Product Prices"],
            ["Alpha Medical", "Sterile gloves", "₹125"],
            ["Beta Medical", "Wheelchair", "₹5,000"],
        ]),
    ])
    _, _, base_records, _, _ = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(),
        empty_target,
        "Company Products and Prices archi sorted.xlsx",
        allow_empty=True,
        source_label="Workbook import",
    )

    merged, added, manual_preview, review_preview, excluded = merge_catalogue_price_workbook_bytes(
        base_records, manual_source, "populated_vendor_prices.xlsx"
    )

    assert base_records.empty
    assert added == 2
    assert len(manual_preview) == 2
    assert review_preview.empty
    assert excluded == 0
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert [(row["Company Name"], row["Product Name"]) for row in products] == [
        ("Alpha Medical", "Sterile gloves"),
        ("Beta Medical", "Wheelchair"),
    ]
    workbook.close()


def test_generated_target_rows_are_preserved_when_manual_source_is_merged():
    generated_records = pd.DataFrame([{
        "company_name": "Existing Vendor",
        "record_type": "catalogue",
        "product_name": "Surgical forceps",
        "sku": "EX-1",
    }])
    generated_target = catalogue_price_workbook_bytes(generated_records)
    manual_source = manual_workbook([
        ("Manual prices", [
            ["Company Name", "Product Name", "Price"],
            ["New Vendor", "Surgical forceps", 250],
        ]),
    ])
    _, _, base_records, _, _ = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(),
        generated_target,
        "Company_Wise_Catalogue_and_Prices.xlsx",
        allow_empty=True,
        source_label="Workbook import",
    )

    merged, added, _, _, _ = merge_catalogue_price_workbook_bytes(
        base_records, manual_source, "manual.xlsx"
    )

    assert added == 1
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert [(row["Company Name"], row["Product Name"]) for row in products] == [
        ("Existing Vendor", "Surgical forceps"),
        ("New Vendor", "Surgical forceps"),
    ]
    workbook.close()


def test_products_prices_sheet_is_preferred_over_needs_review_sheet():
    workbook_bytes = manual_workbook([
        ("Products & Prices", [
            ["Company Name", "Product Name", "Price"],
            ["Alpha Medical", "Sterile gloves", 125],
        ]),
        ("Needs Review", [
            ["Company Name", "Product Name", "Issue"],
            ["", "Unverified row", "Needs review"],
        ]),
    ])

    merged, added, manual_preview, review_preview, excluded = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), workbook_bytes, "Company Products and Prices.xlsx"
    )

    assert added == 1
    assert review_preview.empty
    assert excluded == 0
    assert manual_preview.iloc[0]["product_name"] == "Sterile gloves"
    workbook = load_workbook(BytesIO(merged), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    headers = rows[3]
    products = [dict(zip(headers, row)) for row in rows[4:]]
    assert [row["Product Name"] for row in products] == ["Sterile gloves"]
    workbook.close()


def test_allow_empty_does_not_ignore_unsupported_workbook():
    unsupported = manual_workbook([
        ("Products", [
            ["Product Name", "Price"],
            ["Medical examination gloves", 125],
        ]),
    ])

    with pytest.raises(ValueError, match="Company Name and Product Name"):
        merge_catalogue_price_workbook_bytes(
            pd.DataFrame(), unsupported, "unsupported.xlsx", allow_empty=True
        )


def test_manual_workbook_size_is_limited(monkeypatch):
    assert MAX_MANUAL_CATALOGUE_BYTES == 2 * 1024 * 1024 * 1024
    assert MAX_MANUAL_CATALOGUE_ROWS == 200_000
    monkeypatch.setattr("exports.MAX_MANUAL_CATALOGUE_BYTES", 10)
    content = b"x" * 11
    with pytest.raises(ValueError, match="2 GB"):
        merge_catalogue_price_workbook_bytes(pd.DataFrame(), content, "manual.xlsx")


def test_manual_workbook_row_limit_is_enforced(monkeypatch):
    monkeypatch.setattr("exports.MAX_MANUAL_CATALOGUE_ROWS", 1)
    manual = manual_workbook([
        ("Products & Prices", [
            ["Company Name", "Product Name", "Price"],
            ["Alpha Medical", "Gloves", 125],
            ["Beta Medical", "Wheelchair", 5000],
        ]),
    ])

    with pytest.raises(ValueError, match="200,000 scanned product-row limit"):
        merge_catalogue_price_workbook_bytes(pd.DataFrame(), manual, "manual.xlsx")


def test_rows_missing_company_or_product_are_not_imported():
    manual = manual_workbook(
        [
            (
                "Manual",
                [
                    ["Company Name", "Product Name", "Price"],
                    ["", "Gloves", 100],
                    ["Alpha Medical", "", 100],
                ],
            ),
        ]
    )

    _, added, preview, review, _ = merge_catalogue_price_workbook_bytes(
    pd.DataFrame(), manual, "manual.xlsx"
    )
    assert added == 0
    assert preview.empty
    assert len(review) == 2


def test_description_fallback_and_medicines_are_classified_separately():
    manual = manual_workbook([
        ("Vendor Product Prices", [
            ["Vendor", "Description", "Price"],
            ["Alpha Medical", "Surgical forceps", 125],
            ["Alpha Medical", "Amoxicillin 500 mg capsules", 25],
            ["Alpha Medical", "Office chair", 50],
        ]),
    ])

    _, added, preview, review, excluded = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), manual, "manual.xlsx"
    )

    assert added == 1
    assert preview.iloc[0]["product_name"] == "Surgical forceps"
    assert len(review) == 1
    assert excluded == 1


def test_catalogue_listing_does_not_require_a_price():
    manual = manual_workbook([
        ("Products & Prices", [
            ["Company Name", "Record Type", "Product Name", "Price"],
            ["Alpha Medical", "Catalog listing", "Medical ventilator", None],
        ]),
    ])

    _, added, preview, review, excluded = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), manual, "Company Products and Prices.xlsx"
    )

    assert added == 1
    assert preview.iloc[0]["record_type"] == "catalogue"
    assert review.empty
    assert excluded == 0


def test_source_review_flags_keep_rows_out_of_the_catalogue():
    manual = manual_workbook([
        ("Products & Prices", [
            ["Company Name", "Record Type", "Product Name", "Price", "Review Flag"],
            ["Alpha Medical", "Catalog listing", "Medical ventilator", None, "Needs review"],
        ]),
    ])

    _, added, preview, review, excluded = merge_catalogue_price_workbook_bytes(
        pd.DataFrame(), manual, "Company Products and Prices.xlsx"
    )

    assert added == 0
    assert preview.empty
    assert len(review) == 1
    assert review.iloc[0]["classification_reason"] == "The source workbook flagged this row for review."
    assert excluded == 0


def catalogue_page(tmp_path, monkeypatch):
    from streamlit.testing.v1 import AppTest

    dashboard = Path(__file__).resolve().parents[1] / "vendor_dashboard"
    monkeypatch.syspath_prepend(str(dashboard))
    monkeypatch.setenv("VENDOR_DATA_DIR", str(tmp_path))
    for name in ("DATABASE_URL", "APP_PASSWORD", "CLOUD_DEPLOYMENT"):
        monkeypatch.setenv(name, "")

    page = AppTest.from_file(str(dashboard / "app.py")).run(timeout=40)
    page.radio(key="page").set_value("Catalogue & prices").run()

    assert not page.exception
    return page


def manual_upload(rows=None):
    content = manual_workbook([("Manual prices", [
        ["Company Name", "Product Name", "SKU", "Price", "Currency", "Review Flag"],
        *(rows or [["Beta Medical", "Medical examination gloves", "G1", 100, "INR", ""]]),
    ])])
    return ("manual.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def seed_extracted_catalogue(tmp_path):
    from storage import Store
    from vendor_core import Decision

    store = Store(tmp_path)
    content = manual_workbook([("Catalogue", [
        ["Product Name", "SKU", "Price"], ["Wheelchair", "W1", 5000],
    ])])
    store.save_document("catalogue.xlsx", content, Decision("Alpha Medical", ("Catalogue",), "Test", ""))
    document_id = store.document_id_for("catalogue.xlsx", content)
    store.save_product_records(document_id, [{
        "record_type": "catalogue", "product_name": "Wheelchair", "sku": "W1", "price": 5000,
    }])
    return store


def exported_products(store):
    workbook = load_workbook(BytesIO(catalogue_price_workbook_bytes(store.product_records())), data_only=True)
    rows = list(workbook["All catalogue & prices"].values)
    products = [dict(zip(rows[3], row)) for row in rows[4:]]
    workbook.close()
    return products


def test_manual_upload_automatically_joins_saved_catalogue_and_standard_export(tmp_path, monkeypatch):
    store = seed_extracted_catalogue(tmp_path)
    page = catalogue_page(tmp_path, monkeypatch)
    signature = store.search_index_signature()
    assert len(page.file_uploader) == 1
    assert not page.file_uploader[0].accept_multiple_files
    upload = manual_upload()
    page.file_uploader(key="catalogue_manual_workbook").set_value(upload).run()

    assert not page.exception and not page.error
    assert page.session_state["manual_catalogue_import_result"]["added"] == 1
    assert store.product_record_count() == 2
    assert set(store.product_records().company_name) == {"Alpha Medical", "Beta Medical"}
    assert store.search_index_signature() != signature
    assert {row["Product Name"] for row in exported_products(store)} == {"Wheelchair", "Medical examination gloves"}
    assert any(metric.label == "Saved catalogue records" and metric.value == "2" for metric in page.metric)
    assert [item.key for item in page.get("download_button")] == ["company_wise_catalogue_price_export"]
    assert not any(item.key == "prepare_manual_catalogue_merge" for item in page.button)
    page.text_input(key="product_search").set_value("G1").run()
    assert not page.exception
    assert any(metric.label == "Saved catalogue records" and metric.value == "1" for metric in page.metric)
    assert store.product_record_count() == 2
    page.file_uploader[0].clear().run()
    page.file_uploader[0].set_value(upload).run()
    assert not page.exception and not page.error
    assert page.session_state["manual_catalogue_import_result"]["added"] == 0
    assert store.product_record_count() == 2


def test_upload_survives_clearing_widget_and_new_session(tmp_path, monkeypatch):
    from storage import Store

    page = catalogue_page(tmp_path, monkeypatch)
    page.file_uploader[0].set_value(manual_upload()).run()
    page.file_uploader[0].clear().run()
    assert not page.exception
    assert "manual_catalogue_import_result" not in page.session_state
    page = catalogue_page(tmp_path, monkeypatch)
    assert Store(tmp_path).product_record_count() == 1
    assert any(metric.label == "Saved catalogue records" and metric.value == "1" for metric in page.metric)


def test_invalid_workbook_does_not_change_catalogue(tmp_path, monkeypatch):
    store = seed_extracted_catalogue(tmp_path)
    page = catalogue_page(tmp_path, monkeypatch)
    page.file_uploader[0].set_value(("broken.xlsx", b"invalid", manual_upload()[2])).run()
    assert not page.exception
    assert any("not a readable .xlsx" in error.value for error in page.error)
    assert store.product_record_count() == 1
    assert len(store.documents()) == 1
    page.file_uploader[0].set_value(manual_upload()).run()
    assert not page.exception and not page.error
    assert store.product_record_count() == 2


def test_interrupted_import_offers_retry_instead_of_silently_showing_no_result(tmp_path, monkeypatch):
    from storage import Store

    page = catalogue_page(tmp_path, monkeypatch)
    page.file_uploader[0].set_value(manual_upload()).run()
    # A disconnected rerun can leave the upload token without a completion result.
    del page.session_state['manual_catalogue_import_result']
    page.run()
    assert any('previous workbook import did not finish' in item.value for item in page.warning)
    page.button(key='retry_manual_catalogue_import').click().run()
    assert not page.exception and not page.error
    assert Store(tmp_path).product_record_count() == 1
    assert page.session_state['manual_catalogue_import_result']['duplicates'] == 1


def test_storage_failure_can_be_retried_without_reuploading(tmp_path, monkeypatch):
    import storage

    original = storage.Store.import_catalogue_workbook
    def fail(*args, **kwargs):
        raise RuntimeError("Synthetic database failure")
    page = catalogue_page(tmp_path, monkeypatch)
    monkeypatch.setattr(storage.Store, "import_catalogue_workbook", fail)
    page.file_uploader[0].set_value(manual_upload()).run()
    assert not page.exception
    assert any("Synthetic database failure" in error.value for error in page.error)
    monkeypatch.setattr(storage.Store, "import_catalogue_workbook", original)
    page.run()
    page.button(key="retry_manual_catalogue_import").click().run()
    assert not page.exception and not page.error
    assert storage.Store(tmp_path).product_record_count() == 1


def test_manual_rows_keep_vendor_assignments_and_original_bytes_across_backup(tmp_path):
    from storage import Store

    store = seed_extracted_catalogue(tmp_path / "source")
    upload = manual_upload([
        ["Beta Medical", "Medical examination gloves", "G1", 100, "INR", ""],
        ["Gamma Medical", "Surgical forceps", "F1", 200, "INR", ""],
    ])
    assert store.import_catalogue_workbook(upload[0], upload[1])["added"] == 2
    repeat = store.import_catalogue_workbook("renamed.xlsx", upload[1])
    assert repeat["added"] == 0 and repeat["duplicates"] == 2
    assert Store(tmp_path / "source").product_record_count() == 3
    documents = store.documents()
    manual_docs = documents[documents.method == "Manual catalogue upload"]
    assert set(manual_docs.company_name) == {"Beta Medical", "Gamma Medical"}
    for row in manual_docs.to_dict("records"):
        assert store.safe_path(row["stored_path"]).read_bytes() == upload[1]
    restored = Store(tmp_path / "restored")
    restored.restore_backup(store.backup_bytes())
    records = restored.product_records()
    assert set(zip(records.company_name, records.product_name)) == {
        ("Alpha Medical", "Wheelchair"), ("Beta Medical", "Medical examination gloves"),
        ("Gamma Medical", "Surgical forceps"),
    }
    restored.restore_backup(store.backup_bytes())
    assert restored.product_record_count() == 3
    assert len(restored.documents()) == 3


def test_full_saved_catalogue_includes_review_rows_and_original_price_text(tmp_path, monkeypatch):
    from exports import read_manual_catalogue_workbook
    from storage import Store

    store = Store(tmp_path)
    upload = manual_upload([
        ['Alpha Medical', 'Medical examination gloves', 'G1', 100, 'INR', ''],
        ['Alpha Medical', 'Medical examination gloves', 'G1', 100, 'INR', ''],
        ['Alpha Medical', 'Surgical forceps', 'F1', '₹4,000 + GST', 'INR', ''],
        ['Alpha Medical', 'Surgical forceps', 'F1', '₹5,000 + GST', 'INR', ''],
        ['Beta Medical', 'Wheelchair', 'W1', 'N/A', 'INR', ''],
        ['Alpha Medical', 'Amoxicillin 500 mg capsules', 'D1', 25, 'INR', ''],
    ])
    store.import_catalogue_workbook(upload[0], upload[1])
    assert store.product_record_count() == 1
    assert store.product_record_count(status='all') == 4
    assert store.product_record_count(status='review') == 3
    assert len(store.product_records(status='all', limit=2)) == 2
    original_read = store.read_bytes
    reads = []
    monkeypatch.setattr(store, 'read_bytes', lambda doc: (reads.append(doc), original_read(doc))[1])
    frame = store.catalogue_export_records()
    assert len(frame) == 4 and len(reads) == 1  # One read for all company copies of the same workbook.
    assert set(frame.loc[frame.catalogue_status == 'review', 'source_price_text']) == {'₹4,000 + GST', '₹5,000 + GST', 'N/A'}
    assert (frame.loc[frame.catalogue_status == 'review', 'review_flag'] == 'Needs review').all()
    payload = catalogue_price_workbook_bytes(frame, include_review_details=True)
    workbook = load_workbook(BytesIO(payload), data_only=False)
    sheet = list(workbook['All catalogue & prices'].values)
    records = [dict(zip(sheet[3], row)) for row in sheet[4:]]
    assert len(records) == 4
    assert sum(row['Review Flag'] == 'Needs review' for row in records) == 3
    assert all(row['Review reason'] for row in records)
    assert all(cell.data_type != 'f' for ws in workbook for row in ws for cell in row)
    workbook.close()
    approved, review, excluded = read_manual_catalogue_workbook(payload, 'full.xlsx')
    assert len(approved) == 1 and len(review) == 3 and excluded == 0
    assert store.read_bytes(store.documents().iloc[0].id) == upload[1]


def test_catalogue_status_filter_shows_all_saved_rows_without_approving_them(tmp_path, monkeypatch):
    from storage import Store

    store = Store(tmp_path)
    store.import_catalogue_workbook(*manual_upload([
        ['Alpha Medical', 'Wheelchair', 'W1', 100, 'INR', ''],
        ['Alpha Medical', 'Surgical forceps', 'F1', 'N/A', 'INR', ''],
    ])[:2])
    page = catalogue_page(tmp_path, monkeypatch)
    assert not page.exception
    assert any(metric.label == 'Saved catalogue records' and metric.value == '2' for metric in page.metric)
    page.selectbox(key='product_status').set_value('Validated only').run()
    assert any(metric.label == 'Saved catalogue records' and metric.value == '1' for metric in page.metric)
    page.selectbox(key='product_status').set_value('Needs review').run()
    assert any(metric.label == 'Saved catalogue records' and metric.value == '1' for metric in page.metric)
    assert store.product_classification_counts()['approved'] == 1


def test_manual_review_drugs_and_multi_price_rows_are_preserved_safely(tmp_path):
    from storage import Store

    store = Store(tmp_path / "source")
    upload = manual_upload([
        ["Beta Medical", "Medical examination gloves", "G1", "100, 200", "INR", ""],
        ["Beta Medical", "Surgical forceps", "F1", 200, "INR", "Needs review"],
        ["", "Wheelchair", "W1", 5000, "INR", ""],
        ["Beta Medical", "Medical ventilator", "V1", "invalid", "INR", ""],
        ["Beta Medical", "Amoxicillin 500 mg capsules", "D1", 25, "INR", ""],
    ])
    result = store.import_catalogue_workbook(upload[0], upload[1])
    assert result == {"added": 2, "review": 3, "excluded": 1, "duplicates": 0}
    assert list(store.product_records().price.sort_values()) == [100, 200]
    review = store.product_review_records()
    assert len(review) == 3
    assert "The source workbook flagged this row for review." in set(review.classification_reason)
    assert store.product_records("Amoxicillin").empty
    restored = Store(tmp_path / "restored")
    restored.restore_backup(store.backup_bytes())
    assert restored.product_record_count() == 2
    assert len(restored.product_review_records()) == 3
    with restored.connection() as db:
        db.execute("DELETE FROM vdd_product_records")
    assert restored.backfill_product_records() == 5
    assert restored.product_record_count() == 2
    assert len(restored.product_review_records()) == 3


def test_manual_import_rolls_back_all_vendors_on_database_failure(tmp_path, monkeypatch):
    store = seed_extracted_catalogue(tmp_path)
    original = store._save_product_records
    def fail_second_vendor(db, document_id, records, **kwargs):
        if records and records[0].get("company_name") == "Gamma Medical":
            raise RuntimeError("Synthetic write failure")
        return original(db, document_id, records, **kwargs)
    monkeypatch.setattr(store, "_save_product_records", fail_second_vendor)
    upload = manual_upload([
        ["Beta Medical", "Medical examination gloves", "G1", 100, "INR", ""],
        ["Gamma Medical", "Surgical forceps", "F1", 200, "INR", ""],
    ])
    with pytest.raises(RuntimeError, match="Synthetic write failure"):
        store.import_catalogue_workbook(upload[0], upload[1])
    assert store.product_record_count() == 1
    assert len(store.documents()) == 1
    with store.connection() as db:
        assert db.execute("SELECT COUNT(*) AS n FROM vdd_vendors").fetchone()["n"] == 1
        assert db.execute("SELECT COUNT(*) AS n FROM vdd_audit WHERE action='Manual catalogue upload'").fetchone()["n"] == 0
    monkeypatch.setattr(store, "_save_product_records", original)
    assert store.import_catalogue_workbook(upload[0], upload[1])["added"] == 2


def test_same_filename_with_new_contents_appends_without_overwriting(tmp_path):
    from storage import Store

    store = Store(tmp_path)
    first = manual_upload()
    second = manual_upload([["Beta Medical", "Medical examination gloves", "G1", 150, "INR", ""]])
    assert store.import_catalogue_workbook(first[0], first[1])["added"] == 1
    assert store.import_catalogue_workbook(second[0], second[1])["added"] == 1
    assert sorted(store.product_records().price.tolist()) == [100, 150]
    assert len(store.documents()) == 2


def test_excel_date_cells_import_and_restore_without_serialization_errors(tmp_path, monkeypatch):
    from datetime import datetime, time

    from storage import Store

    content = manual_workbook([("Products & Prices", [
        ["Company Name", "Product Name", "Price", "Date of Conversation", "SKU", "Review Flag"],
        ["Alpha Medical", "Surgical forceps", 125, datetime(2026, 10, 21), time(9, 30), ""],
        ["Beta Medical", "Medical ventilator", 5000, datetime(2026, 10, 20), "V1", "Needs review"],
    ])])
    page = catalogue_page(tmp_path / "source", monkeypatch)
    page.file_uploader[0].set_value(("dated.xlsx", content, manual_upload()[2])).run()
    assert not page.exception and not page.error
    assert page.session_state["manual_catalogue_import_result"]["added"] == 1
    store = Store(tmp_path / "source")
    assert store.product_record_count() == 1
    assert len(store.product_review_records()) == 1
    assert store.product_records().iloc[0].sku == "09:30:00"
    assert store.import_catalogue_workbook("renamed.xlsx", content)["duplicates"] == 2
    restored = Store(tmp_path / "restored")
    restored.restore_backup(store.backup_bytes())
    assert restored.product_record_count() == 1
    assert len(restored.product_review_records()) == 1


def test_catalogue_reloads_old_import_module_and_cached_store(tmp_path, monkeypatch):
    import storage

    page = catalogue_page(tmp_path, monkeypatch)
    old_class = storage.Store
    def old_import(*args, **kwargs):
        raise TypeError("Object of type Timestamp is not JSON serializable")
    monkeypatch.setattr(old_class, "manual_catalogue_import_version", 0)
    monkeypatch.setattr(old_class, "import_catalogue_workbook", old_import)
    page.file_uploader[0].set_value(manual_upload()).run()
    assert not page.exception and not page.error
    assert storage.Store is not old_class
    assert storage.Store(tmp_path).product_record_count() == 1


def test_previous_failed_upload_retries_when_import_version_changes(tmp_path, monkeypatch):
    import storage

    original = storage.Store.import_catalogue_workbook
    def old_import(*args, **kwargs):
        raise TypeError("Object of type Timestamp is not JSON serializable")
    page = catalogue_page(tmp_path, monkeypatch)
    monkeypatch.setattr(storage.Store, "import_catalogue_workbook", old_import)
    page.file_uploader[0].set_value(manual_upload()).run()
    assert page.error
    # This is the unversioned token retained by the previously deployed app.
    page.session_state["manual_catalogue_import_token"] = page.file_uploader[0].value.file_id
    monkeypatch.setattr(storage.Store, "import_catalogue_workbook", original)
    page.run()
    assert not page.exception and not page.error
    assert storage.Store(tmp_path).product_record_count() == 1
    assert page.session_state["manual_catalogue_import_result"]["added"] == 1


def test_manual_sheets_with_different_optional_columns_fit_storage_types(tmp_path):
    from storage import Store

    content = manual_workbook([
        ("Vendor A", [
            ["Company Name", "Product Name", "Price", "MRP", "GST", "Extraction Confidence"],
            ["Alpha Medical", "Surgical forceps", 125, None, None, "high"],
        ]),
        ("Vendor B", [
            ["Company Name", "Product Name", "Record Type", "Price"],
            ["Beta Medical", "Wheelchair", "catalogue", None],
        ]),
    ])
    store = Store(tmp_path)
    assert store.import_catalogue_workbook("mixed-sheets.xlsx", content)["added"] == 2
    with store.connection() as db:
        rows = db.execute("SELECT typeof(price) AS price_type,typeof(mrp) AS mrp_type,typeof(gst) AS gst_type FROM vdd_product_records").fetchall()
    assert all(row[field] in {"real", "null"} for row in rows for field in ("price_type", "mrp_type", "gst_type"))
