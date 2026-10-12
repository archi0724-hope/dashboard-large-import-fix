from io import BytesIO
import zipfile

from openpyxl import load_workbook

from .test_api_export import client


def test_upload_accepts_more_than_multipart_default_file_count(client, monkeypatch):
    from app.main import st
    state = st()
    calls = []
    def save(filename, stream):
        calls.append(filename)
        assert stream.read() == b"company,product\nAcme,Glove"
        return filename
    monkeypatch.setattr(state, "save_upload_stream", save)
    files = [("files", (f"data-{index}.csv", b"company,product\nAcme,Glove", "text/csv")) for index in range(1001)]
    response = client.post("/api/upload", files=files)
    assert response.status_code == 200
    assert len(response.json()["saved"]) == len(calls) == 1001


def test_upload_and_zip_extraction_have_no_default_byte_limit(client):
    from app.main import st
    state = st()
    content = b"source bytes kept intact"
    name = state.save_upload_stream("data.csv", BytesIO(content))
    assert (state.settings.uploads_dir / name).read_bytes() == content
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr("Acme/data.csv", content)
    state.save_upload_stream("files.zip", BytesIO(archive.getvalue()))
    result = state.extract_upload_archive("files.zip")
    assert result["bytes"] == len(content) and result["skipped"] == 0


def test_excel_overflow_creates_continuation_sheets(settings, tmp_path, monkeypatch):
    from app.export import product_workbook
    from app.pipeline import LocalSource
    from app.config import RuntimeConfig
    from app.services import _company_folder_from_source_path
    monkeypatch.setattr(product_workbook, "EXCEL_DATA_ROWS_PER_SHEET", 2)
    folder = tmp_path / "source"
    folder.mkdir()
    (folder / "products.csv").write_text("Company Name,Product Name,SKU,Price,Currency\n" +
        "\n".join(f"Acme,Product {i},SKU{i},{i+1},INR" for i in range(5)), encoding="utf-8")
    output = tmp_path / "workbook.xlsx"
    result = product_workbook.create_product_workbook(LocalSource(folder), settings, RuntimeConfig(), output,
                                                     _company_folder_from_source_path)
    book = load_workbook(output)
    assert result["products"] == 5
    sheets = [book[name] for name in book.sheetnames if name.startswith("Products & Prices")]
    assert [sheet.title for sheet in sheets] == ["Products & Prices", "Products & Prices 2", "Products & Prices 3"]
    assert sum(sheet.max_row - 1 for sheet in sheets) == 5
