import io
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app import served_workbook as served
from app.export.invoice_workbook import HEADERS
from app.vendor_zip_upload import issue_upload_token


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[1]/'vendor_dashboard'))
    from vendor_dashboard.storage import Store
    path = tmp_path/'store'
    monkeypatch.setenv('APP_PASSWORD', 'worksheet-test-password')
    monkeypatch.setattr(served, 'library', lambda: SimpleNamespace(store=Store(path, initialize=False)))
    app = FastAPI()
    app.include_router(served.router)
    return TestClient(app), {'X-Upload-Token': issue_upload_token('worksheet-test-password')}


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Client Call List'
    sheet.append(HEADERS)
    for index in range(302):
        sheet.append([date(2026, 4, 9), 'HOSP013', '<img src=x onerror=alert(1)> Hospital' if index == 0 else 'Example Hospital',
                      'Original City '+str(index), 'Repeated address', '0012345678', '', 'Product '+str(index)])
    other = book.create_sheet('Client Summary')
    other.append(['Duplicate', 'Duplicate', None, 'Notes'])
    other.append(['first', 'second', None, 'Exact wording'])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_original_upload_order_search_all_cells_sheets_and_exact_download(client):
    c, headers = client
    payload = workbook()
    result = c.post('/api/purchase-invoices/workbooks?filename=served%20data.xlsx', headers=headers, content=payload)
    assert result.status_code == 200, result.text
    ident = result.json()['id']
    metadata_response = c.get('/api/purchase-invoices/workbooks', headers=headers)
    assert metadata_response.headers['cache-control'] == 'no-store'
    versions = metadata_response.json()['items']
    assert versions[0]['worksheets'] == [
        {'name':'Client Call List', 'rows':302, 'columns':31}, {'name':'Client Summary', 'rows':1, 'columns':4}]
    base = '/api/purchase-invoices/workbooks/'+ident
    first = c.get(base+'/rows?sheet=Client%20Call%20List', headers=headers).json()
    assert first['headers'] == list(HEADERS)
    assert first['matched'] == first['total'] == 302
    assert len(first['items']) == 250 and first['items'][0]['source_row'] == 2
    assert first['items'][0]['values'][0] == 'Apr 09, 2026'
    assert first['items'][0]['values'][5] == '0012345678'
    next_page = c.get(base+'/rows?sheet=Client%20Call%20List&offset=250', headers=headers).json()
    assert next_page['items'][0]['source_row'] == 252 and len(next_page['items']) == 52
    found = c.get(base+'/rows?sheet=Client%20Call%20List&q=product%20299', headers=headers).json()
    assert found['matched'] == 1 and found['total'] == 302
    assert found['items'][0]['source_row'] == 301 and found['items'][0]['values'][3] == 'Original City 299'
    assert c.get(base+'/rows?sheet=Client%20Call%20List&q=absent', headers=headers).json()['items'] == []
    second = c.get(base+'/rows?sheet=Client%20Summary', headers=headers).json()
    assert second['headers'] == ['Duplicate', 'Duplicate', '', 'Notes']
    assert second['items'][0]['values'] == ['first', 'second', '', 'Exact wording']
    # Every request constructs a new Store: metadata and bytes survive reopened connections.
    downloaded = c.get(base+'/original.xlsx', headers=headers)
    assert downloaded.content == payload and downloaded.headers['cache-control'] == 'no-store'
    duplicate = c.post('/api/purchase-invoices/workbooks?filename=renamed.xlsx', headers=headers, content=payload).json()
    assert duplicate == {'id':ident, 'duplicate':True}
    assert len(c.get('/api/purchase-invoices/workbooks', headers=headers).json()['items']) == 1
    assert c.get(base+'/rows?sheet=Missing', headers=headers).status_code == 404


def test_original_workbook_auth_limits_and_corrupt_uploads(client, monkeypatch):
    c, headers = client
    for path in ('', '/missing/rows?sheet=Client', '/missing/original.xlsx'):
        assert c.get('/api/purchase-invoices/workbooks'+path).status_code == 401
    assert c.post('/api/purchase-invoices/workbooks?filename=served.xlsx', content=workbook()).status_code == 401
    for name in ('../served.xlsx', 'served.xls', 'served.xlsx.exe'):
        assert c.post('/api/purchase-invoices/workbooks', params={'filename':name}, headers=headers, content=b'x').status_code == 422
    assert c.post('/api/purchase-invoices/workbooks?filename=served.xlsx', headers=headers, content=b'broken').status_code == 422
    assert c.post('/api/purchase-invoices/workbooks?filename=served.xlsx', headers=headers, content=b'').status_code == 422
    monkeypatch.setattr(served, 'MAX_BYTES', 2)
    assert c.post('/api/purchase-invoices/workbooks?filename=served.xlsx', headers=headers, content=b'big').status_code == 413
    assert c.get('/api/purchase-invoices/workbooks', headers=headers).json()['items'] == []


def test_original_workbook_rejects_excessive_columns_and_expansion(client):
    c, headers = client
    book = Workbook()
    book.active.append(['column']*101)
    output = io.BytesIO()
    book.save(output)
    response = c.post('/api/purchase-invoices/workbooks?filename=wide.xlsx', headers=headers, content=output.getvalue())
    assert response.status_code == 422 and '100 columns' in response.json()['detail']
    import zipfile
    output = io.BytesIO()
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('xl/workbook.xml', b'0'*(2*1024**2))
    response = c.post('/api/purchase-invoices/workbooks?filename=bomb.xlsx', headers=headers, content=output.getvalue())
    assert response.status_code == 422 and 'compressed' in response.json()['detail']


def test_new_original_workbook_version_keeps_previous_bytes(client):
    c, headers = client
    first = workbook()
    a = c.post('/api/purchase-invoices/workbooks?filename=served.xlsx', headers=headers, content=first).json()['id']
    book = Workbook()
    book.active.append(['Changed column'])
    output = io.BytesIO()
    book.save(output)
    b = c.post('/api/purchase-invoices/workbooks?filename=updated.xlsx', headers=headers, content=output.getvalue()).json()['id']
    assert a != b
    assert len(c.get('/api/purchase-invoices/workbooks', headers=headers).json()['items']) == 2
    assert c.get('/api/purchase-invoices/workbooks/'+a+'/original.xlsx', headers=headers).content == first
