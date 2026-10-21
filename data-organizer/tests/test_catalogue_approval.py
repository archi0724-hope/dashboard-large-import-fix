from __future__ import annotations

import importlib
from io import BytesIO
from pathlib import Path
import zipfile

import pandas as pd
import pytest

from .test_api_export import client


@pytest.fixture
def vendor_store(tmp_path, monkeypatch):
    vendor_dir = Path(__file__).resolve().parents[1] / "vendor_dashboard"
    monkeypatch.syspath_prepend(str(vendor_dir))
    storage = importlib.import_module("storage")
    vendor_core = importlib.import_module("vendor_core")
    store = storage.Store(tmp_path / "vendor-data")
    store.upsert_vendors(pd.DataFrame({"company_name": ["Acme Supplies"]}))
    vendor = store.vendors().iloc[0]
    yield store, vendor_core.Decision("Acme Supplies", ("Price",), "test", "review required")


def add_import(store, decision, filename, content, product_name, sku, price, source_row):
    added = store.save_document(filename, content, decision, source_batch="approval test")
    assert added
    document_id = store.document_id_for(filename, content, "Acme Supplies")
    result = store.stage_product_records(document_id, [{
        "record_type": "price",
        "product_name": product_name,
        "sku": sku,
        "price": price,
        "currency": "INR",
        "pack_size": "box of 100",
        "unit": "box",
        "source_file": filename,
        "source_row": source_row,
    }])
    return result


def test_catalogue_rows_are_not_searchable_until_approved_and_price_changes_are_audited(vendor_store):
    store, decision = vendor_store
    first = add_import(store, decision, "prices.csv", b"first catalogue", "Surgical gloves", "SG-100", 125, 2)
    assert first["staged"] == 1
    assert store.product_records("SG-100").empty

    pending = store.product_import_rows(first["import_id"])[0]
    approved = store.decide_product_import_row(pending["id"], "approve")
    assert approved["status"] == "approved"
    assert store.product_records("SG-100")["price"].tolist() == [125]

    duplicate = store.stage_product_records(
        store.document_id_for("prices.csv", b"first catalogue", "Acme Supplies"),
        [{"record_type": "price", "product_name": "Surgical gloves", "sku": "SG-100",
          "price": 125, "currency": "INR", "pack_size": "box of 100", "unit": "box", "source_row": 2}],
    )
    assert duplicate["staged"] == 0

    revised = add_import(store, decision, "prices-revised.csv", b"revised catalogue",
                         "Surgical gloves", "SG-100", 140, 2)
    revised_row = store.product_import_rows(revised["import_id"])[0]
    assert revised_row["existing_match"]["price"] == 125
    store.edit_product_import_row(revised_row["id"], {"price": "140"})
    result = store.decide_product_import_row(revised_row["id"], "approve")

    assert result["status"] == "approved"
    assert store.product_records("SG-100")["price"].tolist() == [140]
    history = store.product_price_history(product_record_id=result["product_record_id"])
    assert [(entry["old_price"], entry["new_price"], entry["source_file"]) for entry in history] == [
        (125, 140, "prices-revised.csv")
    ]


def test_api_upload_requires_existing_vendor_and_returns_rows_to_review(client, tmp_path, monkeypatch):
    from app import main
    main.catalogue_store.cache_clear()
    monkeypatch.setenv("VENDOR_DATA_DIR", str(tmp_path / "api-vendor-data"))
    monkeypatch.delenv("DATABASE_URL", raising=False)
    store = main.catalogue_store()
    main.catalogue_store.cache_clear()

    store.upsert_vendors(pd.DataFrame({"company_name": ["Acme Supplies"]}))
    vendor_id = store.vendors().iloc[0].canonical_id
    monkeypatch.setattr(main, "catalogue_store", lambda: store)

    invalid_vendor = client.post(
        "/api/catalogue/imports",
        data={"vendor_id": "unknown"},
        files={"file": ("price-list.csv", b"Product Name,SKU,Price\nGloves,G1,125\n", "text/csv")},
    )
    assert invalid_vendor.status_code == 422

    response = client.post(
        "/api/catalogue/imports",
        data={"vendor_id": vendor_id},
        files={"file": ("price-list.csv", b"Product Name,SKU,Price\nGloves,G1,125\n", "text/csv")},
    )
    assert response.status_code == 200
    assert response.json()["records_extracted"] == 1
    assert client.get("/api/catalogue/imports").json()["items"][0]["pending"] == 1
    assert store.product_records("G1").empty

    rows = client.get(f"/api/catalogue/imports/{response.json()['import_id']}/rows").json()["items"]
    published = client.post(f"/api/catalogue/import-rows/{rows[0]['id']}/decide", json={"action": "approve"})
    assert published.status_code == 200
    assert store.product_records("G1")["price"].tolist() == [125]

    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Product Name", "SKU", "Price", "Unit"])
    sheet.append(["Examination gloves", "EG2", 50, "box"])
    excel_file = BytesIO()
    workbook.save(excel_file)
    excel_response = client.post(
        "/api/catalogue/imports",
        data={"vendor_id": vendor_id},
        files={"file": ("price-list.xlsx", excel_file.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert excel_response.status_code == 200
    excel_rows = store.product_import_rows(excel_response.json()["import_id"])
    assert excel_rows[0]["record"]["product_name"] == "Examination gloves"
    assert excel_rows[0]["record"]["price"] == 50


def test_whatsapp_extraction_only_pairs_one_price_with_same_message(monkeypatch):
    vendor_dir = Path(__file__).resolve().parents[1] / "vendor_dashboard"
    monkeypatch.syspath_prepend(str(vendor_dir))
    from catalogue_extraction import extract_whatsapp_records

    records, skipped = extract_whatsapp_records(
        "chat.txt",
        (
            "[17/10/24, 9:10 pm] Devansh: Surgical gloves price ₹1,250 per box\n"
            "[17/10/24, 9:11 pm] Devansh: Gloves ₹100 or ₹200 depending on pack\n"
            "[17/10/24, 9:12 pm] Devansh: We can deliver tomorrow\n"
            "[17/10/24, 9:13 pm] Devansh: ₹500\n"
            "[17/10/24, 9:14 pm] Devansh: Gauze pads 55 INR\n"
        ).encode("utf-8"),
    )

    assert len(records) == 2
    assert skipped == 2
    assert records[0]["product_name"] == "Surgical gloves per box"
    assert records[0]["price"] == 1250
    assert records[0]["source_sheet"] == "Devansh"
    assert records[0]["source_page"] == "17/10/24, 9:10 pm"
    assert records[0]["extraction_confidence"] == "medium"
    assert (records[1]["product_name"], records[1]["price"]) == ("Gauze pads", 55)


def test_whatsapp_zip_upload_stages_rows_and_publishes_only_after_approval(client, tmp_path, monkeypatch):
    from app import main

    main.catalogue_store.cache_clear()
    monkeypatch.setenv("VENDOR_DATA_DIR", str(tmp_path / "whatsapp-vendor-data"))
    monkeypatch.delenv("DATABASE_URL", raising=False)
    store = main.catalogue_store()
    main.catalogue_store.cache_clear()
    store.upsert_vendors(pd.DataFrame({"company_name": ["Acme Supplies"]}))
    vendor_id = store.vendors().iloc[0].canonical_id
    monkeypatch.setattr(main, "catalogue_store", lambda: store)

    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as bundle:
        bundle.writestr(
            "WhatsApp Chat.txt",
            "[17/10/24, 9:10 pm] Acme: Surgical gloves price Rs 125 per box\n",
        )
    response = client.post(
        "/api/catalogue/imports",
        data={"vendor_id": vendor_id},
        files={"file": ("vendor-chat.zip", archive.getvalue(), "application/zip")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["records_extracted"] == 1
    assert store.product_records("Surgical gloves").empty
    rows = store.product_import_rows(payload["import_id"])
    assert rows[0]["record"]["price"] == 125
    assert rows[0]["record"]["source_sheet"] == "WhatsApp Chat.txt"
    assert rows[0]["record"]["source_page"] == "17/10/24, 9:10 pm"
    assert client.post(
        f"/api/catalogue/import-rows/{rows[0]['id']}/decide", json={"action": "approve"}
    ).status_code == 200
    assert store.product_records("Surgical gloves")["price"].tolist() == [125]


def test_approved_catalogue_can_be_searched_and_downloaded_as_sorted_excel(client, tmp_path, monkeypatch):
    from app import main

    main.catalogue_store.cache_clear()
    monkeypatch.setenv("VENDOR_DATA_DIR", str(tmp_path / "catalogue-export-data"))
    monkeypatch.delenv("DATABASE_URL", raising=False)
    store = main.catalogue_store()
    main.catalogue_store.cache_clear()
    store.upsert_vendors(pd.DataFrame({"company_name": ["Acme Supplies"]}))
    vendor_id = store.vendors().iloc[0].canonical_id
    monkeypatch.setattr(main, "catalogue_store", lambda: store)
    upload = client.post(
        "/api/catalogue/imports",
        data={"vendor_id": vendor_id},
        files={"file": ("products.csv", b"Product Name,SKU,Price\nZinc oxide,G2,40\nAlcohol swabs,G1,10\n", "text/csv")},
    )
    assert upload.status_code == 200
    rows = store.product_import_rows(upload.json()["import_id"])
    for row in rows:
        assert client.post(
            f"/api/catalogue/import-rows/{row['id']}/decide", json={"action": "approve"}
        ).status_code == 200

    response = client.get("/api/catalogue/products", params={"vendor_id": vendor_id})
    assert response.status_code == 200
    assert [item["product_name"] for item in response.json()["items"]] == ["Alcohol swabs", "Zinc oxide"]
    assert response.json()["total"] == 2

    export = client.get("/api/catalogue/products/export.xlsx", params={"vendor_id": vendor_id})
    assert export.status_code == 200
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
    sheet = workbook["Approved catalogue"]
    assert [sheet.cell(row=row, column=2).value for row in (2, 3)] == ["Alcohol swabs", "Zinc oxide"]


def test_manual_vendor_workbook_upload_publishes_and_updates_shared_catalogue(vendor_store, monkeypatch):
    store, _ = vendor_store
    vendor_dir = Path(__file__).resolve().parents[1] / "vendor_dashboard"
    monkeypatch.syspath_prepend(str(vendor_dir))
    from import_service import import_catalogue_workbook

    def workbook(price):
        output = BytesIO()
        pd.DataFrame([{
            "Product Name": "Surgical gloves",
            "SKU": "SG-100",
            "Price": price,
            "Unit": "box",
            "Pack Size": "box of 100",
        }]).to_excel(output, index=False)
        return output.getvalue()

    original_workbook = workbook(125)
    revised_workbook = workbook(145)
    first = import_catalogue_workbook(store, "Acme Supplies", "vendor-prices.xlsx", original_workbook)
    assert first["published"] == 1
    assert store.product_records("SG-100")["price"].tolist() == [125]
    assert store.product_import_rows(first["import_id"])[0]["status"] == "approved"

    revised = import_catalogue_workbook(store, "Acme Supplies", "vendor-prices-revised.xlsx", revised_workbook)
    assert revised["published"] == 1
    assert revised["pending_review"] == 0
    records = store.product_records("SG-100")
    assert records["price"].tolist() == [145]
    history = store.product_price_history(product_record_id=int(records.iloc[0]["id"]))
    assert [(item["old_price"], item["new_price"]) for item in history] == [(125, 145)]

    duplicate = import_catalogue_workbook(store, "Acme Supplies", "vendor-prices-revised.xlsx", revised_workbook)
    assert duplicate["duplicate_upload"] is True
    assert store.product_records("SG-100")["price"].tolist() == [145]


def test_vendor_catalogue_page_exposes_manual_workbook_upload(tmp_path, monkeypatch):
    vendor_dir = Path(__file__).resolve().parents[1] / "vendor_dashboard"
    monkeypatch.syspath_prepend(str(vendor_dir))
    monkeypatch.setenv("VENDOR_DATA_DIR", str(tmp_path / "ui-vendor-data"))
    for name in ("DATABASE_URL", "APP_PASSWORD", "CLOUD_DEPLOYMENT"):
        monkeypatch.setenv(name, "")
    storage = importlib.import_module("storage")
    store = storage.Store(tmp_path / "ui-vendor-data")
    store.upsert_vendors(pd.DataFrame({"company_name": ["Acme Supplies"]}))

    from streamlit.testing.v1 import AppTest

    page = AppTest.from_file(str(vendor_dir / "app.py")).run(timeout=30)
    assert not page.exception
    page.radio(key="page").set_value("Catalogue & prices").run(timeout=30)

    assert not page.exception
    assert len(page.file_uploader) == 1
    assert page.file_uploader(key="catalogue_manual_workbook").allowed_type == [".xlsx"]
    assert not any(item.key == "manual_catalogue_upload" for item in page.button)
