"""Formatted, formula-safe Excel and CSV exports."""
from __future__ import annotations

import math
import re
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
from catalogue_policy import clean_and_classify_product

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.exceptions import InvalidFileException
    from openpyxl.workbook.properties import CalcProperties
    OPENPYXL_IMPORT_ERROR = None
except ModuleNotFoundError as error:
    Workbook = None
    load_workbook = None
    InvalidFileException = None
    OPENPYXL_IMPORT_ERROR = error

from vendor_core import DOCUMENT_TYPES, supporting_category

STREAMING_CATALOGUE_EXPORT = True

CATALOGUE_PRICE_COLUMN_OPTIONS = {
    "company_name": ("Company Name", 34),
    "record_type": ("Record Type", 15),
    "product_name": ("Product Name", 38),
    "sku": ("SKU", 18),
    "brand": ("Brand", 22),
    "specification": ("Specification", 42),
    "pack_size": ("Pack Size", 18),
    "unit": ("Unit", 14),
    "price": ("Price", 15),
    "mrp": ("MRP", 15),
    "gst": ("GST", 12),
    "currency": ("Currency", 13),
    "source_file": ("Source File", 42),
    "source_sheet": ("Source Sheet", 20),
    "source_page": ("Source Page", 13),
    "extraction_confidence": ("Extraction Confidence", 22),
    "conversation_date": ("Date of Conversation", 22),
}
MAX_MANUAL_CATALOGUE_BYTES = 2 * 1024 * 1024 * 1024
MAX_MANUAL_CATALOGUE_ROWS = 200_000


def csv_bytes(frame: pd.DataFrame) -> bytes:
    def safe(value):
        if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
            return "'" + value
        return value
    return frame.map(safe).to_csv(index=False).encode("utf-8-sig")


def _export_cell_text(value: object) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, int) or (math.isfinite(value) and value.is_integer()):
            return str(int(value))
    return ILLEGAL_CHARACTERS_RE.sub("", str(value))


def write_text(cell, value):
    cell.value = "" if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))
    cell.data_type = "s"  # A vendor/file name must never become an Excel formula.


def style_sheet(sheet, headers: list[str], widths: list[int], title: str, note: str):
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    write_text(sheet.cell(1, 1), title)
    sheet.cell(1, 1).font = Font(name="Calibri", size=18, bold=True, color="17221B")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.cell(2, 1, note).font = Font(name="Calibri", size=10, color="667268")
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[2].height = 34
    for i, (header, width) in enumerate(zip(headers, widths), 1):
        cell = sheet.cell(4, i, header)
        cell.fill = PatternFill("solid", fgColor="233B2E")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[4].height = 32
    sheet.freeze_panes = "B5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:4"


def finish(sheet):
    sheet.auto_filter.ref = f"A4:{get_column_letter(sheet.max_column)}{max(sheet.max_row,4)}"
    for row in sheet.iter_rows(min_row=5):
        height = 24
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color="000000" if cell.data_type == "f" else "24613C")
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if cell.column == 1 else "center")
            if cell.value == "Yes":
                cell.fill = PatternFill("solid", fgColor="E3F0D8")
            elif cell.value == "No":
                cell.fill = PatternFill("solid", fgColor="FBE9E7")
            elif cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F5F7F2")
            if isinstance(cell.value, str) and cell.data_type != "f":
                width = sheet.column_dimensions[cell.column_letter].width or 15
                height = max(height, min(90, 15 * (math.ceil(len(cell.value) / max(width - 2, 1)) + 1)))
        sheet.row_dimensions[row[0].row].height = height


def workbook_bytes(checklist: pd.DataFrame, documents: pd.DataFrame, per_company: bool = False, aliases: pd.DataFrame | None = None) -> bytes:
    if OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError("Excel export requires the openpyxl package. Add openpyxl to requirements.txt and redeploy the app.") from OPENPYXL_IMPORT_ERROR
    book = Workbook()
    book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    sheet = book.active
    sheet.title = "Document Checklist"
    headers = ["Company Name", *DOCUMENT_TYPES, "Available", "Missing", "Completion", "Files", "Needs review"]
    available_column = 2 + len(DOCUMENT_TYPES)
    missing_column = available_column + 1
    completion_column = missing_column + 1
    note = "Yes = an available file is classified to this category. No = no classified file found. This is NOT validity/compliance verification."
    style_sheet(sheet, headers, [44] + [17] * len(DOCUMENT_TYPES) + [13, 13, 15, 12, 16], (str(checklist.iloc[0]["Company Name"]) + " - DOCUMENT CHECKLIST") if len(checklist)==1 else "VENDOR DOCUMENT CHECKLIST", note)
    for r, record in enumerate(checklist.to_dict("records"), 5):
        for c, header in enumerate(headers, 1):
            value = record[header]
            if header in ("Company Name", *DOCUMENT_TYPES):
                write_text(sheet.cell(r,c), value)
            else:
                sheet.cell(r,c, value)
        first_type = get_column_letter(2)
        last_type = get_column_letter(1 + len(DOCUMENT_TYPES))
        sheet.cell(r, available_column, f'=COUNTIF({first_type}{r}:{last_type}{r},"Yes")')
        sheet.cell(r, missing_column, f'=COLUMNS({first_type}{r}:{last_type}{r})-{get_column_letter(available_column)}{r}')
        sheet.cell(r, completion_column, f'=IFERROR({get_column_letter(available_column)}{r}/COLUMNS({first_type}{r}:{last_type}{r}),0)').number_format = "0%"
        sheet.cell(r,1).comment = Comment("Source: company folder or uploaded vendor master. See Document Register for source paths and classification evidence.", "Vendor Document Desk")
    finish(sheet)
    register = book.create_sheet("Document Register")
    register_headers = ["Company Name", "Document Type", "Filename", "File available", "Review", "Classification", "Source path", "Upload batch", "Uploaded at", "Reason", "Drive original", "Drive processing status"]
    style_sheet(register, register_headers, [40, 25, 72, 16, 16, 25, 90, 42, 28, 82, 65, 24], "DOCUMENT REGISTER", "Uploaded files and Drive references share this register. Drive originals stay in Drive; ZIP member paths locate documents inside archives.")
    for r, doc in enumerate(documents.itertuples(index=False), 5):
        values = [doc.company_name or "Unassigned", ", ".join(doc.types) or supporting_category(doc.filename) or "Other / Needs review", doc.filename,
                  "Yes" if doc.available else "No", "Needs review" if doc.needs_review or not doc.available else "Supporting document" if not doc.types else "Classified",
                  doc.method, doc.original_path, doc.source_batch, doc.uploaded_at, doc.reason, getattr(doc, 'drive_url', ''), getattr(doc, 'drive_status', '')]
        for c, value in enumerate(values,1):
            write_text(register.cell(r,c), value)
    finish(register)
    master = book.create_sheet("Company Master", 0)
    master_headers = ["Canonical ID", "Canonical Company Name", "Match Status"]
    style_sheet(master, master_headers, [18, 52, 24], "COMPANY MASTER", "One canonical row per company. IDs remain stable for future uploads.")
    for r, record in enumerate(checklist.to_dict("records"), 5):
        values = [record.get("canonical_id", ""), record["Company Name"], "Detected from Drive" if str(record.get("canonical_id", "")).startswith('Drive-') else "Confirmed"]
        for c, value in enumerate(values, 1):
            write_text(master.cell(r, c), value)
    finish(master)
    alias_sheet = book.create_sheet("Alias Mapping", 1)
    alias_headers = ["Canonical ID", "Canonical Company", "Original / Detected Name", "Source", "Confidence", "Status"]
    style_sheet(alias_sheet, alias_headers, [18, 42, 42, 42, 14, 24], "COMPANY ALIAS MAPPING", "Aliases are retained for future matching. Uncertain matches are not silently merged.")
    alias_rows = aliases.to_dict("records") if aliases is not None and not aliases.empty else []
    canonical_ids = {record["company_key"]: record.get("canonical_id", "") for record in checklist.to_dict("records")}
    canonical_names = {record["company_key"]: record["Company Name"] for record in checklist.to_dict("records")}
    for r, alias in enumerate(alias_rows, 5):
        values = [canonical_ids.get(alias["canonical_key"], ""), canonical_names.get(alias["canonical_key"], ""),
                  alias["alias_name"], alias["source"], f'{int(alias["confidence"])}%', alias["status"]]
        for c, value in enumerate(values, 1):
            write_text(alias_sheet.cell(r, c), value)
    finish(alias_sheet)
    if per_company:
        names_used = {name.casefold() for name in book.sheetnames}
        for record in checklist.to_dict("records"):
            company = record["Company Name"]
            base = re.sub(r"[\\/*?:\[\]]", " ", company).strip(" '")[:31] or "Company"
            name, suffix = base, 2
            while name.casefold() in names_used:
                tail = f" ({suffix})"; name = base[:31-len(tail)] + tail; suffix += 1
            names_used.add(name.casefold())
            tab = book.create_sheet(name)
            style_sheet(tab, ["Document Type", "Available", "Stored files"], [26, 16, 95], company, note)
            group = documents[documents.company_key == record["company_key"]]
            for r, category in enumerate(DOCUMENT_TYPES,5):
                names = [doc.filename for doc in group.itertuples(index=False) if category in doc.types and doc.available]
                for c, value in enumerate([category, record[category], "\n".join(names) or "No classified file available"],1):
                    write_text(tab.cell(r,c), value)
            other = group[group.types.map(lambda value: not value)]
            for r, doc in enumerate(other.itertuples(index=False), 5+len(DOCUMENT_TYPES)):
                values = [supporting_category(doc.filename) or "Other / Needs review", "Yes" if doc.available else "No", doc.filename]
                for c, value in enumerate(values, 1):
                    write_text(tab.cell(r,c), value)
            finish(tab)
    readme = book.create_sheet("Read me")
    notes = [
        ["Meaning", "Yes means file presence, not authenticity, expiry, completeness or legal compliance."],
        ["No", "No matching available file was found. Check Needs review before concluding the document was never supplied."],
        ["Company ownership", "Company folders take precedence over names inside file names. Generic handover folders are ignored."],
        ["Udyam", "Includes Udyam, Udhyam, MSME and Udyog Aadhaar business registration; not personal Aadhaar."],
        ["ISO", "Recognized ISO certificate names are counted as ISO, regardless of vendor name."],
        ["MD", "MD-form/medical-device licence labels; long certificate serials beginning MD are not MD licences."],
        ["Other files", "Recognised supporting documents appear in the register and company sheet. Only unclear/unavailable files need review. All files are retained."],
        ["Head count", "One row per unique company. Files counts stored document records, not document types. Repeated identical file bytes within one company count once."],
        ["Privacy", "This export may contain sensitive file names. Share only with authorized people."],
        ["Canonical IDs", "Company Master assigns one stable COMP-### ID per company. Alias Mapping preserves detected names and source context."],
        ["Uncertain matches", "Possible matches remain separate and should be corrected through Review files after identifying evidence is supplied."],
    ]
    style_sheet(readme,["Item","Definition"],[26,115],"HOW TO READ THIS WORKBOOK","Source: uploaded vendor master and uploaded document files; no external data or third-party classification API.")
    for r, values in enumerate(notes,5):
        for c, value in enumerate(values,1):write_text(readme.cell(r,c),value)
    finish(readme)
    output = BytesIO(); book.save(output)
    return output.getvalue()


def catalogue_price_workbook_bytes(records: pd.DataFrame, columns: list[str] | None = None) -> bytes:
    """Create a read-only, company-wise workbook from extracted product records.

    This deliberately only reads the extracted rows.  It does not move, rename,
    or delete the original Catalogue/Price documents kept with the checklist.
    """
    if OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError("Excel export requires the openpyxl package. Add openpyxl to requirements.txt and redeploy the app.") from OPENPYXL_IMPORT_ERROR

    source_columns = columns or list(CATALOGUE_PRICE_COLUMN_OPTIONS)
    source_columns = [column for column in source_columns if column in CATALOGUE_PRICE_COLUMN_OPTIONS]
    if not source_columns:
        raise ValueError("Choose at least one Excel column.")
    headers = [CATALOGUE_PRICE_COLUMN_OPTIONS[column][0] for column in source_columns]
    widths = [CATALOGUE_PRICE_COLUMN_OPTIONS[column][1] for column in source_columns]
    frame = records.reindex(columns=list(dict.fromkeys(["company_name", *source_columns])), fill_value="").copy()
    frame["company_name"] = frame["company_name"].fillna("Unassigned").replace("", "Unassigned")
    sort_columns = [column for column in ("company_name", "record_type", "product_name") if column in frame]
    frame = frame.sort_values(sort_columns, kind="stable", na_position="last")

    import xlsxwriter

    output = BytesIO()
    book = xlsxwriter.Workbook(output, {"constant_memory": True, "strings_to_formulas": False, "strings_to_urls": False})
    title_format = book.add_format({"font_name": "Calibri", "font_size": 18, "bold": True, "font_color": "#17221B"})
    note_format = book.add_format({"font_size": 10, "font_color": "#667268", "text_wrap": True})
    header_format = book.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#233B2E", "text_wrap": True})
    row_formats = [book.add_format({"font_name": "Calibri", "font_size": 11, "font_color": "#24613C", "text_wrap": True, "valign": "vcenter", **({"bg_color": "#F5F7F2"} if striped else {})}) for striped in (False, True)]
    names_used: set[str] = set()

    def sheet_name(company: str) -> str:
        base = re.sub(r"[\\/*?:\[\]]", " ", company).strip(" '")[:31] or "Company"
        name, suffix = base, 2
        while name.casefold() in names_used:
            tail = f" ({suffix})"
            name = base[:31-len(tail)] + tail
            suffix += 1
        names_used.add(name.casefold())
        return name

    def add_rows(sheet, company: str, subset: pd.DataFrame) -> None:
        sheet.hide_gridlines(2)
        sheet.freeze_panes(4, 1)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.repeat_rows(0, 3)
        for column, width in enumerate(widths):
            sheet.set_column(column, column, width)
        sheet.set_row(0, 32)
        sheet.merge_range(0, 0, 0, len(headers)-1, company, title_format) if len(headers) > 1 else sheet.write_string(0, 0, company, title_format)
        sheet.set_row(1, 34)
        note = "Extracted catalogue and price rows. Original source files remain stored in the checklist and backups."
        sheet.merge_range(1, 0, 1, len(headers)-1, note, note_format) if len(headers) > 1 else sheet.write_string(1, 0, note, note_format)
        sheet.set_row(3, 32)
        sheet.write_row(3, 0, headers, header_format)
        for row_number, record in enumerate(subset[source_columns].itertuples(index=False, name=None), 4):
            values = [_export_cell_text(value) for value in record]
            sheet.set_row(row_number, 24)
            sheet.write_row(row_number, 0, values, row_formats[row_number % 2])
        sheet.autofilter(3, 0, max(3, len(subset)+3), len(headers)-1)

    all_rows = book.add_worksheet("All catalogue & prices")
    names_used.add("all catalogue & prices")
    add_rows(all_rows, "ALL COMPANIES — CATALOGUE & PRICES", frame)
    for company, subset in frame.groupby("company_name", sort=True, dropna=False):
        add_rows(book.add_worksheet(sheet_name(str(company))), f"{company} — CATALOGUE & PRICES", subset)

    book.close()
    return output.getvalue()


def read_manual_catalogue_workbook(
    manual_workbook: bytes,
    filename: str,
    *,
    allow_empty: bool = False,
    source_label: str = "Manual upload",
) -> tuple[pd.DataFrame, pd.DataFrame, int]:
    """Validate manual product rows without changing the uploaded workbook."""
    if OPENPYXL_IMPORT_ERROR is not None:
        raise RuntimeError("Excel merge requires the openpyxl package.") from OPENPYXL_IMPORT_ERROR
    if Path(filename).suffix.casefold() != ".xlsx":
        raise ValueError("Upload an Excel workbook in .xlsx format.")
    if not manual_workbook:
        raise ValueError("The uploaded Excel workbook is empty.")
    if len(manual_workbook) > MAX_MANUAL_CATALOGUE_BYTES:
        raise ValueError("The uploaded Excel workbook exceeds the 2 GB limit.")

    try:
        workbook = load_workbook(BytesIO(manual_workbook), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise ValueError("This file is not a readable .xlsx workbook.") from error

    header_aliases = {
        field: {re.sub(r"[^a-z0-9]+", "", label.casefold())}
        for field, (label, _) in CATALOGUE_PRICE_COLUMN_OPTIONS.items()
    }
    header_aliases.update({
        "company_name": header_aliases["company_name"] | {
            "company", "vendor", "vendorname", "supplier", "suppliername",
            "hospital", "hospitalname", "organization", "organizationname",
            "organisation", "organisationname", "firm", "firmname", "enterprise",
        },
        "product_name": header_aliases["product_name"] | {
            "product", "products", "productnames", "item", "itemname", "productdescription",
            "itemdescription", "particular", "particulars",
        },
        "sku": header_aliases["sku"] | {"code", "itemcode", "productcode", "model"},
        "price": header_aliases["price"] | {
            "rate", "amount", "cost", "unitprice", "dealerprice", "dealerrate",
            "sellingprice", "salerate", "netprice", "netrate", "rateeach", "productprices",
        },
        "conversation_date": {"dateofconversation", "conversationdate"},
        "mrp": header_aliases["mrp"] | {"maximumretailprice"},
        "gst": header_aliases["gst"] | {"gstpercent", "tax"},
        "currency": header_aliases["currency"] | {"currencycode"},
        "record_type": header_aliases["record_type"] | {"type"},
        "description": {"notes", "sourcepricetext"},
        "review_flag": {"reviewflag"},
        "price_status": {"pricestatus"},
    })
    fields_by_header = {
        alias: field
        for field, aliases in header_aliases.items()
        for alias in aliases
    }
    valid_sheets = []
    preferred_sheet_names = {"allcatalogueprices", "productsprices", "vendorproductprices"}
    preferred_sheet = next(
        (
            sheet for sheet in workbook.worksheets
            if re.sub(r"[^a-z0-9]+", "", sheet.title.casefold()) in preferred_sheet_names
        ),
        None,
    )
    sheets = [preferred_sheet] if preferred_sheet is not None else workbook.worksheets
    manual_rows = []
    review_rows = []
    excluded_rows = 0
    manual_source = f"{source_label}: {Path(filename).name}"
    scanned_rows = 0
    nonempty_data_rows = 0
    recognized_empty_generated_workbook = False

    try:
        for sheet in sheets:
            header_row = None
            fields = {}
            for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), 1):
                found = {
                    fields_by_header[re.sub(r"[^a-z0-9]+", "", str(value).casefold())]: column_number
                    for column_number, value in enumerate(values)
                    if value is not None
                    and re.sub(r"[^a-z0-9]+", "", str(value).casefold()) in fields_by_header
                }
                description_column = next(
                    (
                        column_number
                        for column_number, value in enumerate(values)
                        if value is not None
                        and re.sub(r"[^a-z0-9]+", "", str(value).casefold()) == "description"
                    ),
                    None,
                )
                if "product_name" not in found and description_column is not None:
                    found["product_name"] = description_column
                elif "product_name" in found and description_column is not None:
                    found["description"] = description_column
                if "product_name" in found:
                    header_row = row_number
                    fields = found
                    break
            if header_row is None:
                continue
            normalized_sheet_name = re.sub(r"[^a-z0-9]+", "", sheet.title.casefold())
            if (
                normalized_sheet_name == "allcatalogueprices"
                and "company_name" in fields
                and "product_name" in fields
            ):
                recognized_empty_generated_workbook = True
            price_sheet = "price" in normalized_sheet_name
            if "company_name" not in fields:
                inferred_company = sheet.title.strip()
                generic_sheet_names = {
                    "sheet", "products", "productsprices", "catalogue", "catalog",
                    "pricelist", "prices", "companycontacts", "needsreview",
                    "fileprocessing", "readme", "allcatalogueprices",
                }
                if re.sub(r"[^a-z0-9]+", "", inferred_company.casefold()) in generic_sheet_names:
                    continue
                fields["company_name"] = -1
            valid_sheets.append(sheet.title)
            if sheet.max_row - header_row > MAX_MANUAL_CATALOGUE_ROWS:
                raise ValueError("The workbook exceeds the 200,000 scanned product-row limit.")
            for row_number, values in enumerate(
                sheet.iter_rows(min_row=header_row + 1, values_only=True),
                header_row + 1,
            ):
                scanned_rows += 1
                if scanned_rows > MAX_MANUAL_CATALOGUE_ROWS:
                    raise ValueError("The workbook exceeds the 200,000 scanned product-row limit.")
                row = {
                    field: (
                        sheet.title
                        if column_number == -1
                        else values[column_number] if column_number < len(values) else None
                    )
                    for field, column_number in fields.items()
                }
                if not any(value is not None and str(value).strip() for value in row.values()):
                    continue
                nonempty_data_rows += 1
                company = str(row.get("company_name") or "").strip()
                product = str(row.get("product_name") or "").strip()
                record = {
                    field: "" if value is None else value
                    for field, value in row.items()
                }
                record["company_name"] = company
                record["product_name"] = product
                record_type = re.sub(
                    r"[^a-z]+", " ", str(record.get("record_type") or "").casefold()
                ).strip()
                if record_type in {"catalogue", "catalog", "catalog listing", "product", "product listing"}:
                    record["record_type"] = "catalogue"
                elif record_type == "price" or "price" in record_type:
                    record["record_type"] = "price"
                else:
                    record["record_type"] = (
                        "price" if price_sheet or record.get("price") not in (None, "") else "catalogue"
                    )
                record["source_file"] = manual_source
                record["source_sheet"] = sheet.title
                record["source_page"] = str(row_number)
                if not company:
                    review_rows.append({
                        **record,
                        "classification_reason": "Company/vendor name is missing.",
                    })
                    continue
                if product and not any(character.isalnum() for character in product):
                    continue
                if not any(character.isalnum() for character in product):
                    review_rows.append({
                        **record,
                        "classification_reason": "Product name is empty or not usable.",
                    })
                    continue
                review_markers = ("needs review", "review required", "invalid", "unverified")
                source_flags = " ".join(
                    str(record.get(field) or "").casefold()
                    for field in ("review_flag", "price_status")
                )
                if any(marker in source_flags for marker in review_markers):
                    review_rows.append({
                        **record,
                        "classification_reason": "The source workbook flagged this row for review.",
                    })
                    continue
                record, status, reason = clean_and_classify_product(record, allow_price_lists=True)
                record["classification_reason"] = reason
                if status == "excluded":
                    excluded_rows += 1
                    continue
                if status == "review":
                    review_rows.append(record)
                    continue
                manual_rows.append(record)
                if len(manual_rows) > MAX_MANUAL_CATALOGUE_ROWS:
                    raise ValueError("The workbook exceeds the 200,000 scanned product-row limit.")
    finally:
        workbook.close()

    if allow_empty and recognized_empty_generated_workbook and nonempty_data_rows == 0:
        return pd.DataFrame(), pd.DataFrame(), 0
    if not valid_sheets:
        raise ValueError(
            "No product rows were found. Include Company Name and Product Name columns "
            "(or aliases such as Vendor/Supplier and Product/Item/Description). Include a "
            "company-specific worksheet tab name when there is no company column."
        )
    if not manual_rows:
        if nonempty_data_rows == 0:
            raise ValueError(
                "The worksheet contains column headers but no product rows. "
                "Upload the populated workbook or add product rows below the headers."
            )
        if not review_rows and not excluded_rows:
            raise ValueError("No manual product rows were found. Add Company Name and Product Name to at least one row.")
    manual_frame = pd.DataFrame(manual_rows)
    review_frame = pd.DataFrame(review_rows)
    return manual_frame, review_frame, excluded_rows


def merge_catalogue_price_workbook_bytes(
    records: pd.DataFrame,
    manual_workbook: bytes,
    filename: str,
    *,
    allow_empty: bool = False,
    source_label: str = "Manual upload",
) -> tuple[bytes, int, pd.DataFrame, pd.DataFrame, int]:
    """Build a combined download without modifying the saved catalogue."""
    manual_frame, review_frame, excluded_rows = read_manual_catalogue_workbook(
        manual_workbook, filename, allow_empty=allow_empty, source_label=source_label,
    )
    combined = pd.concat([records, manual_frame], ignore_index=True, sort=False)
    return (
        catalogue_price_workbook_bytes(combined),
        len(manual_frame),
        manual_frame,
        review_frame,
        excluded_rows,
    )
