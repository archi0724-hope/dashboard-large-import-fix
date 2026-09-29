"""Generate a small, deterministic, SYNTHETIC demo data set covering every supported file type.

    python -m scripts.generate_sample_data [--out data/demo_input] [--large 20000]

Names of hospitals/companies are fictional except the "Sawai Man Singh Hospital" example from the specification;
all phone numbers, e-mail addresses, websites and GSTINs are made up. The set deliberately contains the messy
situations the pipeline must cope with (title rows, merged cells, hidden sheet, different headers per sheet,
semicolon CSV in cp1252, multi-page PDF table with header on page 1 only, DOCX table + labelled paragraphs,
free-text notes, a scanned image, a corrupt workbook, an exact duplicate file and ambiguous names).
"""
from __future__ import annotations

import argparse
import csv
import random
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.cleaning.validators import gstin_checksum_char  # noqa: E402


def gstin(state_code: str, pan: str) -> str:
    first14 = f"{state_code}{pan}1Z"
    return first14 + gstin_checksum_char(first14)


G_ARAVALI = gstin("08", "AABCA1111A")
G_MARUDHAR = gstin("08", "AABCM2222B")
G_RAJPUTANA = gstin("08", "AABCR3333C")
G_SARVODAYA = gstin("06", "AABCS4444D")
G_THAR = gstin("08", "AABCT5555E")

ADDR_SMS = "JLN Marg, Jaipur 302004"


def make_excel_rajasthan(path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hospitals"
    ws["A1"] = "HOSPITAL DIRECTORY - RAJASTHAN (compiled 2026)"
    ws.merge_cells("A1:I1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Source: internal survey - draft"
    headers = ["Hospital Name", "Address", "City", "District", "State", "Phone", "Email", "Website", "Registration No"]
    ws.append([])
    ws.append(headers)
    rows = [
        ["SMS HOSPITAL JAIPUR", ADDR_SMS, "Jaipur", "Jaipur", "Rajasthan", "0141-2560101", "", "", ""],
        ["S.M.S. Hospital Jaipur", "Jawaharlal Nehru Marg, Jaipur", "JAIPUR", "Jaipur", "RJ", "0141 2560101", "", "", ""],
        ["Aravali Care Hospital Pvt. Ltd.", "12 Malviya Nagar, Jaipur 302017", "Jaipur", "Jaipur", "Rajasthan", "9876500011",
         "info@aravalicare.example.org", "www.aravalicare.example.org", G_ARAVALI],
        ["ARAVALI CARE HOSP PVT LTD", "12 Malviya Nagar", "Jaipur City", "Jaipur", "Rajasthan", "98765 00011", "", "", G_ARAVALI],
        ["Marudhar Multispeciality Hospital", "Ratanada Road, Jodhpur 342011", "Jodhpur", "Jodhpur", "Rajasthan", "9876500022", "", "", G_MARUDHAR],
        ["ABC Hospital Jaipur", "45 Tonk Road, Jaipur", "Jaipur", "Jaipur", "Rajasthan", "9876500031", "", "", ""],
        ["City Hospital", "Station Road, Jaipur", "Jaipur", "Jaipur", "Rajasthan", "9876500041", "", "", ""],
        [None] * 9,
        ["Shri Ganesh Nursing Home", "Rangbari Road, Kota 324007", "Kota", "Kota", "Rajasthan", "9876500051", "", "", ""],
        ["Total", "", "", "", "", "", "", "", ""],
    ]
    for r in rows:
        ws.append(r)

    ws2 = wb.create_sheet("Clinics")
    ws2.append(["Name of Institution", "Location", "Town", "Contact No.", "GSTIN", "Contact Person"])
    clinic_rows = [
        ["Sawai Man Singh Hospital", "Jawaharlal Nehru Marg, Jaipur 302004", "Jaipur", 9876500061, "", "Dr. R. Sharma"],
        ["Thar Pharma Traders", "Sardarpura, Jodhpur", "Jodhpur", 9876500071, G_THAR, "Mr. K. Singh"],
        ["Marudhar Multi Speciality Hospital", "Ratanada, Jodhpur", None, 9876500022, "", ""],
        ["Dr. Kapoor Hospital", "MI Road, Jaipur", "Jaipur", 9876500081, "", ""],
        ["Doctor Kapoor Hospital", "M.I. Road, Jaipur", "Jaipur", 9876500081, "", ""],
    ]
    for r in clinic_rows:
        ws2.append(r)
    ws2.merge_cells("C3:C4")       # merged 'Town' cells: Jodhpur applies to two rows

    ws3 = wb.create_sheet("Archive")
    ws3.append(["Hospital Name", "City", "State"])
    ws3.append(["Old Town Dispensary", "Ajmer", "Rajasthan"])
    ws3.sheet_state = "hidden"
    wb.save(path)


def make_excel_kota(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kota"
    ws.append(["Organization", "Address", "Dist", "State", "Mobile", "Category"])
    for r in [
        ["ABC Hospital Kota", "Nayapura, Kota", "Kota", "Rajasthan", "9876500091", "Private"],
        ["Sri Ganesh Nursing Home Kota", "Rangbari Road", "Kota", "Rajasthan", "9876500051", "Private"],
        ["St. Mary Hospital", "Talwandi, Kota", "Kota", "Rajasthan", "9876500101", "Private"],
        ["Saint Mary Hospital Kota", "Talwandi, Kota", "Kota", "Rajasthan", "", "Private"],
        ["ABC Hospital", "", "", "", "", ""],
    ]:
        ws.append(r)
    wb.save(path)


def make_xls(path: Path) -> None:
    try:
        import xlwt
    except ImportError:
        print("xlwt not installed - skipping legacy .xls sample")
        return
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Directory")
    for c, h in enumerate(["Institution Name", "City", "State", "Ph No"]):
        ws.write(0, c, h)
    for r, row in enumerate([
        ["Gurugram City Hospital", "Gurugram", "Haryana", "9876500111"],
        ["Gurgaon City Hospital", "Gurgaon", "Haryana", "9876500111"],
        ["Sarvodaya Medical Centre", "Faridabad", "Haryana", "9876500121"],
        ["City Hospital", "Ambala", "Haryana", "9876500131"],
    ], start=1):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    wb.save(str(path))


def make_csv(path: Path) -> None:
    rows = [
        ["Company Name", "Address", "City", "State", "GST No", "Mobile", "E-mail"],
        ["Rajputana Surgicals Private Limited", "Sitapura Industrial Area", "Jaipur", "Rajasthan", G_RAJPUTANA, "9876500141", "sales@rajputana.example.org"],
        ["Rajputana Surgicals Pvt Ltd", "Sitapura, Jaipur", "Jaipur", "Rajasthan", G_RAJPUTANA, "98765 00141", ""],
        ["M/s Rajputana Surgicals", "", "Jaipur", "Rajasthan", "", "", ""],
        ["Café Médical Supplies Co.", "Vaishali Nagar", "Jaipur", "Rajasthan", "", "98765", "not-an-email@@x"],
        ["Sarvodaya Medical Center", "Sector 8", "Faridabad", "Haryana", G_SARVODAYA, "9876500121", ""],
    ]
    with open(path, "w", encoding="cp1252", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerows(rows)


def make_pdf(path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4)
    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 8)])
    t1 = Table([
        ["Hospital Name", "City", "State", "Phone"],
        ["Sawai Man Singh Hospital", "Jaipur", "Rajasthan", "0141-2560101"],
        ["Aravali Care Hospital", "Jaipur", "Rajasthan", "9876500011"],
        ["Eastern Star Hospital", "Jaipur", "Rajasthan", "9876500151"],
    ])
    t1.setStyle(grid)
    t2 = Table([   # continuation on page 2: NO header row (must be carried over)
        ["Marudhar Multispeciality Hospital", "Jodhpur", "Rajasthan", "9876500022"],
        ["Western Star Hospital", "Jaipur", "Rajasthan", "9876500161"],
    ])
    t2.setStyle(grid)
    t3 = Table([
        ["Hospital Name", "City", "State", "Phone"],
        ["Gurugram City Hospital", "Gurugram", "Haryana", "9876500111"],
    ])
    t3.setStyle(grid)
    doc.build([
        Paragraph("Empanelled hospitals - annex", styles["Heading1"]), t1, PageBreak(),
        t2, Spacer(1, 12),
        Paragraph("Note: the list above is indicative and subject to verification by the district office.", styles["Normal"]),
        PageBreak(), t3,
    ])


def make_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Demo Author"
    d.add_heading("Hospitals of Haryana", level=1)
    d.add_paragraph("The following facilities were surveyed in 2026.")
    t = d.add_table(rows=1, cols=3)
    for i, h in enumerate(["Name of Hospital", "City", "Phone"]):
        t.rows[0].cells[i].text = h
    for row in [
        ("Medicare Plus Hospital", "Gurugram", "9876500171"),
        ("Medicare Plus Hospital Pvt. Ltd.", "Gurgaon", "9876500171"),
        ("Sarvodaya Medical Centre", "Faridabad", "9876500121"),
    ]:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
    d.add_heading("Additional facilities", level=2)
    for block in [
        ("Name: Panipat Eye Care Hospital", "Address: GT Road, Panipat", "City: Panipat", "State: Haryana", "Phone: 9876500181"),
        ("Name: Karnal Heart Institute", "Address: Sector 12, Karnal", "City: Karnal", "State: Haryana", "Email: info@karnalheart.example.org"),
    ]:
        for line in block:
            d.add_paragraph(line)
    d.save(path)


def make_txt(path: Path) -> None:
    lines = [
        "FIELD NOTES - JAIPUR VISIT",
        "",
        "1. SMS Hosp., Jaipur | 0141-2560101 | www.smshospital.example.org",
        "2. Aravali Care Hospital, Malviya Nagar, Jaipur - 302017, Ph: 9876500011",
        "3. Eastern Star Hospital, Jaipur, Rajasthan",
        "Met the medical superintendent on Tuesday; the trust intends to expand the outpatient wing next year and asked us to share the list.",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def make_scanned_png(path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont

    fonts = ["/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans.ttf"]
    font = None
    for f in fonts:
        if Path(f).exists():
            font = ImageFont.truetype(f, 34)
            break
    font = font or ImageFont.load_default(size=34)
    img = Image.new("RGB", (1500, 520), "white")
    draw = ImageDraw.Draw(img)
    draw.text((40, 30), "Hospital List (scanned)", fill="black", font=font)
    y = 110
    for line in ["Eastern Star Hospital, Jaipur, Rajasthan", "Kota Care Clinic, Kota, Rajasthan", "Panipat Eye Care Hospital, Panipat, Haryana"]:
        draw.text((40, y), line, fill="black", font=font)
        y += 90
    img = img.rotate(0.6, expand=False, fillcolor="white")
    img.save(path)


def make_large(path: Path, n_records: int, n_entities: int, seed: int = 7) -> None:
    """Benchmark data: many noisy variants of a pool of synthetic entities (CSV)."""
    rnd = random.Random(seed)
    first = ["Sri", "Shri", "Maha", "Ananta", "Kalyan", "Nav", "Jeevan", "Arogya", "Sanjeevani", "Lakshmi", "Bharat", "Vijay",
             "Rajdhani", "Aravali", "Thar", "Marudhar", "Mewar", "Hada", "Vindhya", "Gomti", "Yamuna", "Ganga", "Kaveri"]
    second = ["Ram", "Krishna", "Shanti", "Surya", "Chandra", "Tara", "Deep", "Kiran", "Asha", "Prem", "Om", "Sagar", "Nirmal", "Pushp"]
    kinds = ["Hospital", "Nursing Home", "Medical Centre", "Clinic", "Diagnostics"]
    cities = [("Jaipur", "Rajasthan"), ("Jodhpur", "Rajasthan"), ("Kota", "Rajasthan"), ("Udaipur", "Rajasthan"),
              ("Gurugram", "Haryana"), ("Faridabad", "Haryana"), ("Ludhiana", "Punjab"), ("Lucknow", "Uttar Pradesh"),
              ("Indore", "Madhya Pradesh"), ("Pune", "Maharashtra")]
    ents = []
    seen = set()
    while len(ents) < n_entities:
        nm = f"{rnd.choice(first)} {rnd.choice(second)} {rnd.choice(second)} {rnd.choice(kinds)}"
        city = rnd.choice(cities)
        key = (nm, city[0])
        if key in seen:
            continue
        seen.add(key)
        ents.append((nm, city, f"{rnd.randint(1, 250)} {rnd.choice(['MG Road', 'Station Road', 'Civil Lines', 'Sector 4', 'Ring Road'])}",
                     f"9{rnd.randint(100000000, 999999999)}"))

    def vary(name: str) -> str:
        v = name
        if rnd.random() < .3:
            v = v.replace("Hospital", rnd.choice(["Hosp.", "Hospital", "HOSPITAL"]))
        if rnd.random() < .2:
            v = v.upper()
        if rnd.random() < .15:
            v = v.replace("Sri", "Shri")
        if rnd.random() < .1 and len(v) > 8:
            i = rnd.randrange(2, len(v) - 2)
            v = v[:i] + v[i + 1:]           # dropped letter (typo)
        if rnd.random() < .2:
            v = v + f", {rnd.choice(['Jaipur', 'Kota'])}" if rnd.random() < .05 else v
        return v

    with open(path, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["Hospital Name", "Address", "City", "State", "Phone"])
        for _ in range(n_records):
            nm, (city, state), addr, phone = rnd.choice(ents)
            w.writerow([vary(nm), addr if rnd.random() < .8 else "", city if rnd.random() < .9 else "",
                        state if rnd.random() < .8 else "", phone if rnd.random() < .6 else ""])


def generate(out: Path, large: int = 0) -> list[Path]:
    if out.exists():
        shutil.rmtree(out)
    (out / "regional").mkdir(parents=True)
    made = []
    for fn, fnc in [
        ("hospitals_rajasthan.xlsx", make_excel_rajasthan), ("regional/kota_hospitals.xlsx", make_excel_kota),
        ("legacy_directory.xls", make_xls), ("companies.csv", make_csv), ("empanelled_hospitals.pdf", make_pdf),
        ("haryana_facilities.docx", make_docx), ("field_notes.txt", make_txt), ("scanned_directory.png", make_scanned_png),
    ]:
        fnc(out / fn)
        made.append(out / fn)
    shutil.copy(out / "companies.csv", out / "regional" / "companies_copy.csv")           # exact duplicate file
    (out / "corrupt_file.xlsx").write_bytes(b"PK\x03\x04 this is not a real workbook" + bytes(range(64)))
    (out / "readme_notes.md").write_text("# not a supported format\n", encoding="utf-8")
    if large:
        make_large(out / "large_synthetic.csv", large, max(50, large // 8))
    return made


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(ROOT / "data" / "demo_input"))
    ap.add_argument("--large", type=int, default=0, help="also write a synthetic CSV with this many records (benchmark)")
    a = ap.parse_args()
    generate(Path(a.out), a.large)
    print(f"Sample data written to {a.out}")
