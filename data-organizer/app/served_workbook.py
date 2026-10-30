"""Versioned original invoice workbooks and read-only worksheet browsing."""
from __future__ import annotations

import hashlib
import json
import math
import zipfile
from datetime import date, datetime, timezone
from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response
from openpyxl import load_workbook
from starlette.concurrency import run_in_threadpool
from urllib.parse import quote

from .drive.index_api import authorize, library

def private_response(response: Response):
    response.headers['Cache-Control'] = 'no-store'


router = APIRouter(prefix='/api/purchase-invoices/workbooks', dependencies=[Depends(authorize), Depends(private_response)])
MAX_BYTES = 20 * 1024**2
SCHEMA = '''CREATE TABLE IF NOT EXISTS vdd_invoice_workbooks (
    id TEXT PRIMARY KEY, schema_version INTEGER NOT NULL DEFAULT 1,
    filename TEXT NOT NULL, payload BYTEA NOT NULL,
    metadata_json TEXT NOT NULL, created_at TEXT NOT NULL)'''


def database():
    store = library().store
    with store.connection() as db:
        db.execute(SCHEMA)
    return store


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.strftime('%b %d, %Y')
    if isinstance(value, float) and not math.isfinite(value):
        return ''
    return str(value)


def inspect_workbook(payload):
    if not payload or len(payload) > MAX_BYTES:
        raise ValueError('Choose an Excel workbook up to 20 MB.')
    try:
        with zipfile.ZipFile(BytesIO(payload)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(entry.file_size for entry in entries) > 64 * 1024**2:
                raise ValueError('Workbook exceeds the 5,000-entry or 64 MB expanded-size limit.')
            if any(entry.flag_bits & 1 or entry.file_size / max(1, entry.compress_size) > 1000 for entry in entries):
                raise ValueError('Encrypted or excessively compressed workbooks are not supported.')
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        try:
            if len(book.worksheets) > 50:
                raise ValueError('Workbooks support up to 50 worksheets.')
            metadata, total_cells = [], 0
            for sheet in book.worksheets:
                if not sheet.max_row or not sheet.max_column or sheet.max_row > 50001 or sheet.max_column > 100:
                    raise ValueError('Each worksheet supports up to 50,000 data rows and 100 columns with valid dimensions.')
                # Iterate to validate the XML and actual dimensions before accepting original bytes.
                row_count, columns = 0, sheet.max_column
                for values in sheet.iter_rows(values_only=True):
                    row_count += 1
                    total_cells += len(values)
                    if row_count > 50001 or len(values) > 100 or total_cells > 1_000_000:
                        raise ValueError('Workbook exceeds the row, column or 1,000,000-cell processing limit.')
                metadata.append({'name': sheet.title, 'rows': max(0, row_count-1), 'columns': columns})
            return {'schema_version': 1, 'worksheets': metadata}
        finally:
            book.close()
    except ValueError:
        raise
    except Exception as error:  # Parser failures are invalid uploads, not server errors.
        raise ValueError('This file is not a readable .xlsx workbook.') from error


def save_workbook(filename, payload):
    metadata = inspect_workbook(payload)
    ident = hashlib.sha256(payload).hexdigest()
    store = database()
    with store.connection() as db:
        existing = db.execute('SELECT id FROM vdd_invoice_workbooks WHERE id=?', (ident,)).fetchone()
        if existing:
            return {'id': ident, 'duplicate': True}
        if store.cloud and db.execute('SELECT pg_database_size(current_database()) AS n').fetchone()['n'] + len(payload) >= 850 * 1024**2:
            raise HTTPException(507, 'The shared database is nearly full. No workbook was saved.')
        db.execute('INSERT INTO vdd_invoice_workbooks (id,schema_version,filename,payload,metadata_json,created_at) VALUES (?,1,?,?,?,?) ON CONFLICT (id) DO NOTHING',
                   (ident, filename, payload, json.dumps(metadata), datetime.now(timezone.utc).isoformat()))
    return {'id': ident, 'duplicate': False}


@router.post('')
async def upload(request: Request, filename: str):
    if len(filename) > 200 or '/' in filename or '\\' in filename or Path(filename).suffix.lower() != '.xlsx':
        raise HTTPException(422, 'Choose an .xlsx workbook with a valid filename.')
    content = bytearray()
    async for chunk in request.stream():
        content.extend(chunk)
        if len(content) > MAX_BYTES:
            raise HTTPException(413, 'Choose an Excel workbook up to 20 MB.')
    try:
        return await run_in_threadpool(save_workbook, filename, bytes(content))
    except (ValueError, TypeError, IndexError) as error:
        raise HTTPException(422, str(error)) from None


@router.get('')
def workbooks():
    with database().connection() as db:
        rows = db.execute('SELECT id,filename,metadata_json,created_at FROM vdd_invoice_workbooks ORDER BY created_at DESC,id').fetchall()
    return {'items': [{'id': row['id'], 'filename': row['filename'], 'created_at': row['created_at'],
                       **json.loads(row['metadata_json'])} for row in rows]}


def original(ident):
    with database().connection() as db:
        row = db.execute('SELECT filename,payload,metadata_json FROM vdd_invoice_workbooks WHERE id=?', (ident,)).fetchone()
    if not row:
        raise HTTPException(404, 'Original workbook not found.')
    return row


@router.get('/{ident}/original.xlsx')
def download(ident: str):
    row = original(ident)
    return Response(bytes(row['payload']), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Cache-Control': 'no-store', 'Content-Disposition': "attachment; filename*=UTF-8''"+quote(row['filename'], safe='')})


@router.get('/{ident}/rows')
def worksheet_rows(ident: str, sheet: str, q: str = '', offset: int = 0):
    row = original(ident)
    metadata = json.loads(row['metadata_json'])
    info = next((item for item in metadata['worksheets'] if item['name'] == sheet), None)
    if not info:
        raise HTTPException(404, 'Worksheet not found.')
    book = load_workbook(BytesIO(bytes(row['payload'])), read_only=True, data_only=True, keep_links=False)
    try:
        rows = book[sheet].iter_rows(values_only=True)
        headers = [cell_text(value) for value in next(rows, ())]
        offset, matched, items = max(0, offset), 0, []
        query = q[:500].casefold().strip()
        for source_row, values in enumerate(rows, 2):
            values = [cell_text(value) for value in values]
            if query and not any(query in value.casefold() for value in values):
                continue
            if offset <= matched < offset+250:
                items.append({'source_row': source_row, 'values': values})
            matched += 1
        return {'headers': headers, 'items': items, 'total': info['rows'], 'matched': matched,
                'columns': info['columns'], 'offset': offset, 'limit': 250}
    finally:
        book.close()
