"""Index original Drive sources incrementally; only metadata and rows reach SQL."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import tempfile
import threading
import uuid
from contextlib import contextmanager
from pathlib import Path, PurePosixPath

from .index_store import DriveIndex
from .public_drive import MAX_DOCUMENT, download_public_file, list_public_files
from .range_zip import DriveRangeFile, archive_members

OWNER = uuid.uuid4().hex
LOCK = threading.Lock()
CANCEL = threading.Event()
ACTIVE = None
EXTRACTABLE = {'.xlsx','.xls','.csv','.txt','.log','.pdf','.docx','.json'}
COMPANY = {'company','companyname','vendor','vendorname','supplier','suppliername','manufacturer','manufacturername','hospital','hospitalname'}
PRODUCT = {'product','productname','item','itemname','description','particulars','itemdescription','productdescription'}
PRICE = {'price','rate','mrp','unitprice','amount','netprice','offerprice'}


def entry_key(drive_id, member=''):
    return hashlib.sha256((drive_id + '\0' + member).encode()).hexdigest()


def normalize(value):
    return re.sub(r'[^a-z0-9]', '', str(value).casefold())


def table_records(rows, sheet='', fallback_company=''):
    headers = None
    for number, values in enumerate(rows, 1):
        if isinstance(values,dict):
            number,values = values['row'],values['values']
        cells = [str(value).strip() if value is not None else '' for value in values]
        if not any(cells):
            continue
        if headers is None:
            normalized = {normalize(value) for value in cells}
            if normalized & (COMPANY | PRODUCT | PRICE):
                headers = [(cell or f'Column {i+1}') + (f' [{i+1}]' if cell in cells[:i] else '') for i,cell in enumerate(cells)]
                continue
            # Preserve title/preamble rows too; never discard unknown rows.
            yield {'kind':'text','company':fallback_company,'product':'','sheet':sheet,'row':number,'values':cells}
            continue
        data = {headers[i] if i < len(headers) else f'Column {i+1}':value for i,value in enumerate(cells) if value}
        normalized = {normalize(key):value for key,value in data.items()}
        company = next((value for key,value in normalized.items() if key in COMPANY),fallback_company)
        product = next((value for key,value in normalized.items() if key in PRODUCT),'')
        price = next((value for key,value in normalized.items() if key in PRICE),'')
        yield {'kind':'product' if product else 'company' if company else 'row',
               'company':company,'product':product,'price':price,'sheet':sheet,'row':number,'values':data}


def extracted_records(path, fallback_company=''):
    ext = path.suffix.lower()
    if ext == '.xlsx':
        from .xlsx_stream import sheets
        for name,rows in sheets(path):
            yield from table_records(rows,name,fallback_company)
    elif ext == '.xls':
        import xlrd
        book = xlrd.open_workbook(path,on_demand=True)
        try:
            for sheet in book.sheets():
                yield from table_records((sheet.row_values(row) for row in range(sheet.nrows)),sheet.name,fallback_company)
        finally:
            book.release_resources()
    elif ext == '.csv':
        with path.open(encoding='utf-8-sig',errors='replace',newline='') as source:
            yield from table_records(csv.reader(source),fallback_company=fallback_company)
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path,strict=False)
        for number,page in enumerate(reader.pages,1):
            text = page.extract_text() or ''
            if not text.strip():
                yield {'kind':'needs_ocr','company':fallback_company,'product':'','page':number,'text':'Image-only page; OCR needed.'}
            else:
                for offset in range(0,len(text),20000):
                    yield {'kind':'document','company':fallback_company,'product':'','page':number,'part':offset//20000+1,'text':text[offset:offset+20000]}
    elif ext == '.docx':
        from docx import Document
        doc = Document(path)
        for number,paragraph in enumerate(doc.paragraphs,1):
            if paragraph.text.strip():
                yield {'kind':'document','company':fallback_company,'product':'','row':number,'text':paragraph.text}
        for number,table in enumerate(doc.tables,1):
            yield from table_records(([cell.text for cell in row.cells] for row in table.rows),f'Table {number}',fallback_company)
    else:
        with path.open(encoding='utf-8-sig',errors='replace') as source:
            for number,line in enumerate(source,1):
                if line.strip():
                    yield {'kind':'text','company':fallback_company,'product':'','row':number,'text':line.rstrip()}


def safe_member(name):
    path = PurePosixPath(name.replace('\\','/'))
    return not (path.is_absolute() or '..' in path.parts or re.match(r'^[A-Za-z]:',name))


def discover(index, folder):
    files = list_public_files(folder)
    for file in files:
        if CANCEL.is_set():
            return
        key = entry_key(file['id'])
        base = {'entry_key':key,'folder_id':folder,'drive_id':file['id'],'source_name':file['path'],
                'size_bytes':0,'version':'public-file','status':'pending','error':''}
        index.state(folder,'running','Listing '+file['path'],OWNER)
        if file['name'].lower().endswith('.zip'):
            try:
                with DriveRangeFile(file['id']) as remote, archive_members(remote) as archive:
                    entries = []
                    signature = hashlib.sha256()
                    for member in archive.infolist():
                        signature.update(f'{member.filename}|{member.CRC}|{member.file_size}|{member.compress_size}'.encode())
                        if member.is_dir():
                            continue
                        ext = Path(member.filename).suffix.lower()
                        status,reason = 'pending',''
                        if not safe_member(member.filename):
                            status,reason = 'blocked','Unsafe archive member path.'
                        elif member.flag_bits & 1:
                            status,reason = 'blocked','Encrypted ZIP member requires a password.'
                        elif member.file_size > MAX_DOCUMENT:
                            status,reason = 'blocked','Original is indexed, but exceeds the 128 MB document processing limit.'
                        elif ext not in EXTRACTABLE:
                            status,reason = 'needs_ocr' if ext in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'} else 'unsupported','Image needs OCR.' if ext in {'.png','.jpg','.jpeg','.tif','.tiff','.bmp','.webp'} else 'This file format is not supported for text indexing.'
                        entries.append({**base,'entry_key':entry_key(file['id'],member.filename),'member_path':member.filename,
                            'size_bytes':member.file_size,'version':f'{member.CRC}:{member.file_size}:{member.compress_size}',
                            'status':status,'error':reason})
                    index.discover([{**base,'size_bytes':remote.size,'version':signature.hexdigest(),'status':'container'}])
                    for offset in range(0,len(entries),200):
                        index.discover(entries[offset:offset+200])
            except Exception as error:  # noqa: BLE001 - retain failures per source without aborting the folder
                index.discover([{**base,'status':'failed','error':str(error) if isinstance(error,ValueError) else 'Cannot read ZIP directory: '+type(error).__name__}])
        else:
            ext = Path(file['name']).suffix.lower()
            if ext not in EXTRACTABLE or file.get('native'):
                base.update(status='unsupported',error='Native Google documents need Google authorization.' if file.get('native') else 'This format is visible in the library but cannot be text-indexed. Upload an equivalent ZIP/CSV/XLSX for extraction.')
            if not file.get('native'):
                try:
                    with DriveRangeFile(file['id']) as remote:
                        base['size_bytes'] = remote.size
                except (ValueError,OSError):
                    # Listing an inaccessible file must still retain its inventory row.
                    pass
            index.discover([base])
            if base['status'] == 'pending':
                # Public loose files have no reliable version in the folder page.
                # Re-fetch and compare their SHA-256 on each explicit refresh.
                index.mark(key,'pending')


@contextmanager
def source_copy(entry, target):
    """Temporary one-document copy; validate ZIP size and CRC before indexing."""
    if entry['member_path']:
        with DriveRangeFile(entry['drive_id']) as remote, archive_members(remote) as archive:
            member = archive.getinfo(entry['member_path'])
            if not safe_member(member.filename) or member.file_size > MAX_DOCUMENT or member.flag_bits & 1:
                raise ValueError('Unsafe, encrypted or oversized archive member.')
            with archive.open(member) as source, target.open('wb') as output:
                total = 0
                while chunk := source.read(4 * 1024**2):
                    total += len(chunk)
                    if total > MAX_DOCUMENT:
                        raise ValueError('Document exceeds the 128 MB processing limit.')
                    output.write(chunk)
    else:
        download_public_file(entry['drive_id'],target)
    try:
        yield target
    finally:
        target.unlink(missing_ok=True)


def run_index(store, folder):
    index = DriveIndex(store)
    # A session-level lock also prevents a resumed process and old deployment
    # from indexing the same folder concurrently. It releases on disconnect.
    with store.connection() as lease:
        if store.cloud:
            lease.connection.autocommit = True
            key = int(hashlib.sha256(folder.encode()).hexdigest()[:15],16)
            if not lease.execute('SELECT pg_try_advisory_lock(?) AS locked',(key,)).fetchone()['locked']:
                return
        try:
            index.state(folder,'running','Discovering files and ZIP contents',OWNER)
            discover(index,folder)
            for entry in index.pending(folder):
                if CANCEL.is_set():
                    index.state(folder,'paused','Paused. Resume continues from saved files and rows.',OWNER)
                    return
                if not entry['member_path'] and entry['source_name'].lower().endswith('.zip'):
                    continue
                if store.cloud:
                    with store.connection() as db:
                        size = int(db.execute('SELECT pg_database_size(current_database()) AS n').fetchone()['n'])
                    if size > 850 * 1024**2:
                        index.state(folder,'paused','Index paused near the 1 GiB database limit. More index storage is needed to continue.',OWNER)
                        return
                name = entry['member_path'] or entry['source_name']
                index.state(folder,'running','Reading '+name,OWNER)
                index.mark(entry['entry_key'],'reading')
                try:
                    with tempfile.TemporaryDirectory(prefix='drive-index-') as work:
                        target = Path(work) / ('document'+Path(name).suffix.lower())
                        with source_copy(entry,target):
                            if not entry['size_bytes']:
                                index.set_size(entry['entry_key'],target.stat().st_size)
                            with target.open('rb') as content:
                                digest = hashlib.file_digest(content,'sha256').hexdigest()
                            existing = index.existing(entry['entry_key'])
                            if existing['content_hash'] == digest and existing['record_count'] and target.suffix != '.pdf':
                                index.mark(entry['entry_key'],'indexed','',digest,existing['record_count'])
                                continue
                            if existing['content_hash'] and existing['content_hash'] != digest:
                                index.reset_records(entry['entry_key'])
                            total,batch,ocr = 0,[],False
                            for record in extracted_records(target):
                                ocr = ocr or record['kind'] == 'needs_ocr'
                                # Retain traceability and keep unverified rows out of the approved catalogue.
                                record.update(source_name=entry['source_name'],member_path=entry['member_path'],
                                    source_url=f"https://drive.google.com/file/d/{entry['drive_id']}/view",review_status='Needs review')
                                batch.append(record)
                                total += 1
                                if len(batch) == 200:
                                    if any(len(json.dumps(row,default=str)) > 128 * 1024 for row in batch):
                                        raise ValueError('A source row exceeds the 128 KB index-row limit; split that row or document.')
                                    index.save_records(entry,batch)
                                    batch.clear()
                                    index.state(folder,'running',f'Reading {name}: {total:,} rows saved',OWNER)
                                    if total % 2000 == 0 and store.cloud:
                                        with store.connection() as db:
                                            size = int(db.execute('SELECT pg_database_size(current_database()) AS n').fetchone()['n'])
                                        if size > 850 * 1024**2:
                                            index.mark(entry['entry_key'],'pending')
                                            index.state(folder,'paused','Index paused near the 1 GiB storage limit. Saved rows are retained.',OWNER)
                                            return
                                    if CANCEL.is_set():
                                        index.mark(entry['entry_key'],'pending')
                                        index.state(folder,'paused','Paused; saved rows are retained.',OWNER)
                                        return
                            if batch:
                                index.save_records(entry,batch)
                            index.mark(entry['entry_key'],'partial' if ocr else 'indexed',
                                'Some pages need OCR.' if ocr else '',digest,total)
                except Exception as error:  # noqa: BLE001 - retain completed documents and report this file
                    index.mark(entry['entry_key'],'failed',str(error) if isinstance(error,ValueError) else 'Extraction failed: '+type(error).__name__)
            counts = index.status(folder)['counts']
            incomplete = sum(counts.get(status,0) for status in ('failed','blocked','unsupported','needs_ocr','partial','pending','reading'))
            index.state(folder,'finished_with_issues' if incomplete else 'done',
                f'Pass finished. {incomplete:,} entries need attention.' if incomplete else 'All supported documents indexed.',OWNER)
        except Exception as error:  # noqa: BLE001 - persist a resumable job failure
            index.state(folder,'failed','Index interrupted: '+type(error).__name__+'. Resume to retry.',OWNER)
        finally:
            if store.cloud:
                lease.execute('SELECT pg_advisory_unlock(?)',(key,))


def start_index(store, folder):
    global ACTIVE
    with LOCK:
        if ACTIVE and ACTIVE.is_alive():
            return
        CANCEL.clear()
        ACTIVE = threading.Thread(target=run_index,args=(store,folder),name='drive-index',daemon=True)
        ACTIVE.start()
